"""Stage 4 security review: report aggregates, money visibility, certificates, employment,
page caps and money in the audit log."""

import asyncio
from collections.abc import Callable
from io import BytesIO
from typing import Any

import httpx
from fastapi.routing import APIRoute
from openpyxl import load_workbook

from app.main import app
from tests.conftest import SignedIn, branch_id, make_role, sql
from tests.edu import instructor_rights
from tests.money import S, make_payout
from tests.reports import EM, ML, D, G, employ, finished_cohort, student_of, world

MONEY = ("income", "payroll", "ads", "profit", "cac")


def header(r: httpx.Response) -> list[Any]:
    assert r.status_code == 200, r.text
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    return [c.value for c in next(ws.iter_rows(max_row=1))]


# ---- money on the dashboard, master list and graduates -----------------------------------


async def test_dashboard_money_needs_the_money_sections(
    admin: httpx.AsyncClient, browser: Callable[[], httpx.AsyncClient], login_as: SignedIn
) -> None:
    """income ⇐ payments view, payroll ⇐ instructors view, ads/CAC ⇐ marketing view,
    profit ⇐ all three. Hidden = null in JSON and no column in Excel."""
    await world(admin, browser)
    bare = await login_as("bare", await make_role("Только дашборд", {"dashboard": "view"}))
    o = (await bare.get(f"{D}/overview")).json()
    assert all(o[k] is None for k in MONEY), o
    assert o["leads"] == 5  # counts stay
    for m in (await bare.get(f"{D}/by-month")).json()["items"]:
        assert all(m[k] is None for k in ("income", "payroll", "ads", "profit")), m
    for c in (await bare.get(f"{D}/by-course")).json()["items"]:
        assert c["income"] is None
    cols = header(await bare.get(f"{D}/by-month/export"))
    assert not {"Доход", "Зарплата", "Реклама", "Прибыль"} & set(cols), cols

    pay_only = await login_as(
        "payonly", await make_role("Дашборд+оплаты", {"dashboard": "view", "payments": "view"})
    )
    o = (await pay_only.get(f"{D}/overview")).json()
    assert o["income"] == "1600.00"
    assert (o["payroll"], o["ads"], o["profit"], o["cac"]) == (None, None, None, None)

    full = await login_as(
        "full",
        await make_role(
            "Всё",
            {"dashboard": "view", "payments": "view", "instructors": "view", "marketing": "view"},
        ),
    )
    o = (await full.get(f"{D}/overview")).json()
    assert (o["income"], o["payroll"], o["ads"], o["profit"]) == (
        "1600.00", "400.00", "1200.00", "0.00",
    )  # fmt: skip


async def test_dashboard_lead_widgets_need_leads_view(
    admin: httpx.AsyncClient, browser: Callable[[], httpx.AsyncClient], login_as: SignedIn
) -> None:
    await world(admin, browser)
    bare = await login_as("bare", await make_role("Только дашборд", {"dashboard": "view"}))
    o = (await bare.get(f"{D}/overview")).json()
    assert o["recent_leads"] == [] and o["callbacks"]["items"] == []
    sales = await login_as(
        "sales", await make_role("Дашборд+лиды", {"dashboard": "view", "leads": "view"})
    )
    assert len((await sales.get(f"{D}/overview")).json()["recent_leads"]) == 5


async def test_master_list_and_graduates_money_need_payments_view(
    admin: httpx.AsyncClient, browser: Callable[[], httpx.AsyncClient], login_as: SignedIn
) -> None:
    await world(admin, browser)
    viewer = await login_as(
        "ml", await make_role("Отчёты", {"master_list": "view", "graduates": "view"})
    )
    for row in (await viewer.get(ML)).json()["items"]:
        assert row["income"] is None and row["ltv"] is None
    cols = header(await viewer.get(f"{ML}/export"))
    assert "Доход" not in cols and "LTV/чел" not in cols
    grads = (await viewer.get(f"{G}/students")).json()["items"]
    assert grads and all(g["paid"] is None for g in grads)
    assert "Оплачено" not in header(await viewer.get(f"{G}/students/export"))
    # with payments view the numbers are there
    rows = (await admin.get(ML)).json()["items"]
    assert any(r["income"] not in (None, "0.00") for r in rows)


async def test_branch_bound_user_gets_no_other_branch_in_any_grouping(
    admin: httpx.AsyncClient, browser: Callable[[], httpx.AsyncClient], login_as: SignedIn
) -> None:
    await world(admin, browser)
    from tests.money import manager_role

    mgr = await login_as("mgr", await manager_role(), branch="Душанбе")
    own = {"branch_id": str(await branch_id("Душанбе"))}
    other = {"branch_id": str(await branch_id("Истаравшан"))}
    # by-course and by-month equal the admin's Душанбе-only numbers, also with old periods
    for url, params in (
        (f"{D}/by-course", {}),
        (f"{D}/by-course", {"date_from": "2020-01-01"}),
        (f"{D}/by-month", {}),
        (ML, {"page_size": 100}),
        (ML, {"status": "активный", "date_from": "2020-01-01"}),
    ):
        mine = (await mgr.get(url, params=params)).json()["items"]
        theirs = (await admin.get(url, params={**params, **own})).json()["items"]
        assert [r.get("course_name") or r.get("month") for r in mine] == [
            r.get("course_name") or r.get("month") for r in theirs
        ], url
        assert all(r.get("branch_name") in (None, "Душанбе") for r in mine), url
        assert all(r.get("course_name") != "Курс Ist" or r.get("income") is None for r in mine)
    for url in (
        f"{D}/by-month/export",
        f"{ML}/export",
        f"{G}/students/export",
        f"{G}/courses",
        f"{G}/courses/export",
        EM,
        f"{EM}/export",
    ):
        assert (await mgr.get(url, params=other)).status_code == 404, url


async def test_instructor_is_denied_every_report_endpoint_even_with_grants(
    admin: httpx.AsyncClient, browser: Callable[[], httpx.AsyncClient]
) -> None:
    w = await world(admin, browser)
    mine = w["instructor_client"]
    enrollment = w["done_ids"]["Выпускник Один"]
    grad = await student_of(enrollment)
    async with instructor_rights(
        dashboard="change", master_list="change", graduates="change", payments="change"
    ):
        for route in app.routes:
            if not isinstance(route, APIRoute):
                continue
            path = route.path
            if not path.startswith(
                (
                    "/api/v1/dashboard",
                    "/api/v1/master-list",
                    "/api/v1/graduates",
                    "/api/v1/employment",
                )
            ):
                continue
            url = path.replace("{enrollment_id}", enrollment).replace(
                "{employment_id}", "00000000-0000-0000-0000-000000000000"
            )
            for method in route.methods:
                r = await mine.request(
                    method,
                    url,
                    json={"student_id": grad, "organization": "X", "employed_at": "2026-01-01"}
                    if method in ("POST", "PATCH")
                    else None,
                    headers={"X-Record-Version": "1"} if method == "PATCH" else None,
                )
                assert r.status_code == 403, (method, url, r.status_code)


# ---- certificates ------------------------------------------------------------------------


async def test_certificate_needs_change_scope_and_numbers_stay_unique_under_load(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    names = {f"Студент {i}": (40, 40) for i in range(6)}
    _, ids = await finished_cohort(admin, names)
    viewer = await login_as("v", await make_role("Вид", {"graduates": "view"}))
    first = next(iter(ids.values()))
    assert (await viewer.post(f"{G}/{first}/certificate")).status_code == 403
    ist = await login_as(
        "ist", await make_role("Ист", {"graduates": "change"}), branch="Истаравшан"
    )
    assert (await ist.post(f"{G}/{first}/certificate")).status_code == 404
    # everyone at once, twice each: one number per enrollment, all distinct, never overwritten
    results = await asyncio.gather(
        *(admin.post(f"{G}/{e}/certificate") for e in [*ids.values(), *ids.values()])
    )
    ok = [r.json()["certificate_no"] for r in results if r.status_code == 200]
    assert len(ok) == len(ids) == len(set(ok))
    assert sum(r.status_code == 409 for r in results) == len(ids)
    stored = await sql("SELECT certificate_no FROM enrollment WHERE certificate_no IS NOT NULL")
    assert sorted(n for (n,) in stored) == sorted(ok)


# ---- employment --------------------------------------------------------------------------


async def test_employment_scope_by_student_branch_and_audit_sentences(
    admin: httpx.AsyncClient, browser: Callable[[], httpx.AsyncClient], login_as: SignedIn
) -> None:
    w = await world(admin, browser)
    grad = await student_of(w["done_ids"]["Выпускник Один"])
    ist = await login_as(
        "ist", await make_role("Ист", {"graduates": "change"}), branch="Истаравшан"
    )
    assert (await employ(ist, grad)).status_code == 404
    job = (await employ(admin, grad)).json()
    assert (await ist.delete(f"{EM}/{job['id']}")).status_code == 404
    assert (await ist.get(EM, params={"student_id": grad})).json()["total"] == 0
    r = await admin.patch(f"{EM}/{job['id']}", json={"organization": "Новая"})
    assert r.status_code == 200
    assert (await admin.delete(f"{EM}/{job['id']}")).status_code == 204
    log = (await admin.get("/api/v1/audit", params={"section": "graduates"})).json()["items"]
    texts = [i["text"] for i in log if i["action"].startswith("employment.")]
    assert len(texts) == 3
    assert all("Выпускник Один" in t and "None" not in t for t in texts), texts


# ---- hardening: page cap -----------------------------------------------------------------


def paged_routes() -> list[str]:
    """Every GET whose OpenAPI schema takes a `page` query parameter."""
    return [
        path
        for path, ops in app.openapi()["paths"].items()
        if any(p["name"] == "page" for p in ops.get("get", {}).get("parameters", []))
    ]


async def test_every_paged_list_caps_the_page_number(admin: httpx.AsyncClient) -> None:
    routes = paged_routes()
    assert len(routes) >= 10, routes
    for path in routes:
        url = path.replace("{instructor_id}", "00000000-0000-0000-0000-000000000000")
        r = await admin.get(url, params={"page": 10**15})
        assert r.status_code == 422, (path, r.status_code)
        r = await admin.get(url, params={"page": 100_000})
        assert r.status_code in (200, 404), (path, r.status_code)


# ---- hardening: money in the audit log ---------------------------------------------------


async def test_audit_hides_money_from_users_without_the_money_sections(
    admin: httpx.AsyncClient, browser: Callable[[], httpx.AsyncClient], login_as: SignedIn
) -> None:
    w = await world(admin, browser)
    await make_payout(admin, w["live"], amount="777.77")
    r = await admin.patch(
        f"{S}/cohorts/{w['live']['id']}/rate", json={"instructor_rate": "98765.43"}
    )
    assert r.status_code == 200, r.text
    auditor = await login_as("aud", await make_role("Журнал", {"audit": "view"}))
    items = (await auditor.get("/api/v1/audit", params={"page_size": 100})).json()["items"]
    assert items
    entities = {i["entity"] for i in items if i["action"] != "export"}
    assert not {"payment", "instructor_payout"} & entities, entities
    blob = repr(items)
    assert "98765.43" not in blob and "instructor_rate" not in blob and "budget" not in blob
    for q in ("98765.43", "777.77"):
        found = (await auditor.get("/api/v1/audit", params={"q": q})).json()
        assert found["total"] == 0, q
    for section in ("payments", "instructors"):
        found = (await auditor.get("/api/v1/audit", params={"section": section})).json()
        assert found["total"] == 0, section
    xlsx = await auditor.get("/api/v1/audit/export")
    assert xlsx.status_code == 200
    ws = load_workbook(BytesIO(xlsx.content)).active
    assert ws is not None
    cells = " ".join(str(c) for row in ws.iter_rows(values_only=True) for c in row)
    assert "98765.43" not in cells and "777.77" not in cells
    # with the sections the entries are there
    money = await login_as(
        "money",
        await make_role(
            "Журнал+деньги",
            {"audit": "view", "payments": "view", "instructors": "view", "marketing": "view"},
        ),
    )
    found = (await money.get("/api/v1/audit", params={"q": "98765.43"})).json()
    assert found["total"] >= 1
    items = (await money.get("/api/v1/audit", params={"page_size": 100})).json()["items"]
    assert {"payment", "instructor_payout"} <= {i["entity"] for i in items}
