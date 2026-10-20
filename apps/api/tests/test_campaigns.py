"""Stage 3: campaigns and marketing numbers: CPL / CAC, totals, ads in the profit, scope."""

from io import BytesIO

import httpx
from openpyxl import load_workbook

from tests.conftest import SignedIn, branch_id, make_role
from tests.edu import instructor_rights, make_course
from tests.money import enroll, pay, running_cohort
from tests.money import manager_role as money_manager
from tests.sales import CA, L, M, make_campaign, make_lead, manager_role


async def _convert(c: httpx.AsyncClient, lead: dict[str, str]) -> None:
    r = await c.post(f"{L}/{lead['id']}/convert", json={})
    assert r.status_code == 200, r.text


async def test_cpl_and_cac_math_including_zero_leads(admin: httpx.AsyncClient) -> None:
    empty = await make_campaign(admin, "Пустая", budget=1000)
    assert (empty["leads"], empty["converted"], empty["enrolled"]) == (0, 0, 0)
    assert (empty["cpl"], empty["cac"]) == (None, None)  # «—», never a division error
    camp = await make_campaign(admin, "Осень", budget=1000)
    leads = [
        await make_lead(admin, f"Лид {i}", phone=f"90000000{i}", campaign_id=camp["id"])
        for i in range(4)
    ]
    r = (await admin.get(f"{CA}/{camp['id']}")).json()
    assert (r["leads"], r["converted"], r["cpl"], r["cac"]) == (4, 0, "250.00", None)
    await _convert(admin, leads[0])
    await _convert(admin, leads[1])
    r = (await admin.get(f"{CA}/{camp['id']}")).json()
    assert (r["leads"], r["converted"], r["cpl"], r["cac"]) == (4, 2, "250.00", "500.00")
    # money rounds to the cent
    third = await make_campaign(admin, "Три", budget=100)
    for i in range(3):
        await make_lead(admin, f"Т{i}", campaign_id=third["id"])
    assert (await admin.get(f"{CA}/{third['id']}")).json()["cpl"] == "33.33"


async def test_enrolled_counts_converted_leads_with_an_enrollment(
    admin: httpx.AsyncClient,
) -> None:
    cohort = await running_cohort(admin)
    camp = await make_campaign(admin)
    a = await make_lead(admin, "Анна", phone="901000001", campaign_id=camp["id"])
    b = await make_lead(admin, "Борис", phone="902000002", campaign_id=camp["id"])
    await admin.post(f"{L}/{a['id']}/convert", json={"cohort_id": cohort["id"]})
    await _convert(admin, b)  # a student, but not in a cohort yet
    r = (await admin.get(f"{CA}/{camp['id']}")).json()
    assert (r["leads"], r["converted"], r["enrolled"]) == (2, 2, 1)


async def test_list_totals_and_filters(admin: httpx.AsyncClient) -> None:
    a = await make_campaign(admin, "Фейсбук", budget=600, platform="Facebook Ads")
    b = await make_campaign(admin, "Тикток", budget=400, platform="TikTok", status="активна")
    for i in range(3):
        await make_lead(admin, f"A{i}", campaign_id=a["id"], phone=f"90100000{i}")
    lead = await make_lead(admin, "B0", campaign_id=b["id"], phone="902000000")
    await _convert(admin, lead)
    body = (await admin.get(CA)).json()
    assert {i["name"]: i["leads"] for i in body["items"]} == {"Фейсбук": 3, "Тикток": 1}
    t = body["totals"]
    assert (t["campaigns"], t["budget"], t["leads"], t["converted"]) == (2, "1000.00", 4, 1)
    assert (t["cpl"], t["cac"]) == ("250.00", "1000.00")
    only = (await admin.get(CA, params={"status": "активна"})).json()
    assert [i["name"] for i in only["items"]] == ["Тикток"]
    assert only["totals"]["budget"] == "400.00"
    assert (await admin.get(CA, params={"q": "фейс"})).json()["totals"]["campaigns"] == 1
    assert (await admin.get(CA, params={"platform": "Google"})).json()["totals"]["cpl"] is None
    none = (await admin.get(CA, params={"date_from": "2030-01-01"})).json()
    assert (none["items"], none["totals"]["budget"]) == ([], "0.00")


async def test_profit_is_income_minus_ads_minus_payroll(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin, back=1, price=800)
    e = await enroll(admin, cohort)
    await pay(admin, e, amount=800, confirm=True)
    f = (await admin.get(f"{M}/finance")).json()
    assert (f["income"], f["ads"], f["profit"]) == ("800.00", "0.00", "800.00")
    await make_campaign(admin, "Осень", budget=150)
    await make_campaign(admin, "Старая", budget=50, start_date="2025-01-10")
    f = (await admin.get(f"{M}/finance")).json()
    assert (f["income"], f["ads"], f["profit"]) == ("800.00", "200.00", "600.00")
    # ads count in the period the campaign started in
    f = (await admin.get(f"{M}/finance", params={"date_from": "2026-01-01"})).json()
    assert f["ads"] == "150.00"
    f = (await admin.get(f"{M}/finance", params={"date_to": "2025-12-31"})).json()
    assert (f["ads"], f["income"]) == ("50.00", "0.00")


async def test_finance_ads_follow_the_branch_scope(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    await make_campaign(admin, "Душанбе", budget=100)
    await make_campaign(admin, "Истаравшан", branch="Истаравшан", budget=70)
    assert (await admin.get(f"{M}/finance")).json()["ads"] == "170.00"
    anna = await login_as("anna", await money_manager(), branch="Душанбе")
    assert (await anna.get(f"{M}/finance")).json()["ads"] == "100.00"


async def test_campaign_validation(admin: httpx.AsyncClient) -> None:
    branch = str(await branch_id("Душанбе"))
    base = {"name": "К", "branch_id": branch, "platform": "TikTok", "start_date": "2026-09-01"}
    r = await admin.post(CA, json={**base, "end_date": "2026-08-01"})
    assert (r.status_code, r.json()["code"]) == (422, "campaign_dates_invalid")
    assert (await admin.post(CA, json={**base, "budget": "-1"})).status_code == 422
    assert (await admin.post(CA, json={**base, "platform": "Радио"})).status_code == 422
    assert (await admin.post(CA, json={**base, "status": "бред"})).status_code == 422
    assert (await admin.post(CA, json={**base, "name": "  "})).status_code == 422
    assert (await admin.post(CA, json={**base, "budget": "10.123"})).status_code == 422
    nil = "00000000-0000-0000-0000-000000000000"
    assert (await admin.post(CA, json={**base, "course_id": nil})).status_code == 404
    ok = await make_campaign(admin, "Нормальная", end_date="2026-10-01")
    r = await admin.patch(f"{CA}/{ok['id']}", json={"end_date": "2026-01-01"})
    assert r.json()["code"] == "campaign_dates_invalid"
    r = await admin.patch(f"{CA}/{ok['id']}", json={"budget": "0", "end_date": None})
    assert (r.json()["budget"], r.json()["end_date"]) == ("0.00", None)
    assert r.json()["version"] == 2


async def test_campaign_version_conflict_and_delete(admin: httpx.AsyncClient) -> None:
    camp = await make_campaign(admin)
    await admin.patch(f"{CA}/{camp['id']}", json={"budget": "5"}, headers={"If-Match": "1"})
    stale = await admin.patch(f"{CA}/{camp['id']}", json={"budget": "6"}, headers={"If-Match": "1"})
    assert (stale.status_code, stale.json()["code"]) == (409, "stale_version")
    await make_lead(admin, "Анна", campaign_id=camp["id"])
    r = await admin.delete(f"{CA}/{camp['id']}")
    assert (r.status_code, r.json()["code"]) == (409, "campaign_in_use")
    free = await make_campaign(admin, "Пустая")
    assert (await admin.delete(f"{CA}/{free['id']}")).status_code == 204


async def test_lead_and_campaign_must_share_a_branch(admin: httpx.AsyncClient) -> None:
    camp = await make_campaign(admin, branch="Истаравшан")
    r = await admin.post(
        L,
        json={
            "full_name": "Анна",
            "branch_id": str(await branch_id("Душанбе")),
            "campaign_id": camp["id"],
        },
    )
    assert (r.status_code, r.json()["code"]) == (422, "campaign_branch_mismatch")


async def test_scope_rights_and_export(admin: httpx.AsyncClient, login_as: SignedIn) -> None:
    mine = await make_campaign(admin, "Моя", branch="Душанбе", budget=100)
    other = await make_campaign(admin, "Чужая", branch="Истаравшан", budget=900)
    anna = await login_as("anna", await manager_role(), branch="Душанбе")
    names = [i["name"] for i in (await anna.get(CA)).json()["items"]]
    assert names == ["Моя"]
    assert (await anna.get(CA)).json()["totals"]["budget"] == "100.00"
    assert (await anna.get(f"{CA}/{other['id']}")).status_code == 404
    # the manager's starting right on marketing is view
    assert (await anna.patch(f"{CA}/{mine['id']}", json={"budget": "1"})).status_code == 403
    assert (await anna.post(CA, json={"name": "x"})).status_code == 403
    maker = await make_role("Маркетолог", {"marketing": "change"})
    editor = await login_as("ed", maker, branch="Душанбе")
    assert (await editor.patch(f"{CA}/{other['id']}", json={"budget": "1"})).status_code == 404
    assert (await editor.delete(f"{CA}/{other['id']}")).status_code == 404
    r = await editor.post(CA, json={"name": "Н", "platform": "TikTok", "start_date": "2026-10-01"})
    assert r.status_code == 201
    assert r.json()["branch_name"] == "Душанбе"
    nobody = await login_as("nobody", await make_role("Никто", {"leads": "change"}))
    assert (await nobody.get(CA)).status_code == 403
    r = await anna.get(f"{CA}/export")
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][6:12] == ("Бюджет", "Лиды", "Записались", "Зачислено", "CPL", "CAC")
    assert [x[0] for x in rows[1:]] == ["Н", "Моя"]  # newest start first, own branch only


async def test_instructor_has_no_marketing_access(login_as: SignedIn) -> None:
    async with instructor_rights(marketing="change"):
        ins = await login_as("ins", "instructor")
        assert (await ins.get(CA)).status_code == 403
        assert (await ins.get(f"{CA}/export")).status_code == 403
        assert (await ins.post(CA, json={})).status_code == 403


async def test_campaign_of_a_course_and_audit(admin: httpx.AsyncClient) -> None:
    course = await make_course("Python")
    camp = await make_campaign(admin, "Курсовая", course_id=str(course))
    assert camp["course_name"] == "Python"
    await admin.patch(f"{CA}/{camp['id']}", json={"budget": "250", "status": "активна"})
    items = (await admin.get("/api/v1/audit", params={"section": "marketing"})).json()["items"]
    texts = [i["text"] for i in items]
    assert any("создана кампания Курсовая" in t for t in texts)
    assert any("бюджет 1 000 → 250 TJS" in t and "запланирована → активна" in t for t in texts)
