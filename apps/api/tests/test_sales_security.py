"""Stage 3 security regressions: sales input handling, rights and scope."""

from io import BytesIO

import httpx
from openpyxl import load_workbook

from tests.conftest import SignedIn, branch_id, make_role, sql
from tests.edu import make_student
from tests.sales import CA, L, M, csv_bytes, make_campaign, make_lead


async def _csv(c: httpx.AsyncClient, path: str, data: bytes, **params: str) -> httpx.Response:
    return await c.post(
        f"{L}{path}", content=data, params=params, headers={"Content-Type": "text/csv"}
    )


async def test_nul_characters_in_lead_fields_are_a_422_not_a_500(
    admin: httpx.AsyncClient,
) -> None:
    branch = str(await branch_id("Душанбе"))
    for field in ("full_name", "comment", "phone", "refusal_reason"):
        body = {"full_name": "Анна", "branch_id": branch, "status": "отказ", field: "a\x00b"}
        r = await admin.post(L, json=body)
        assert r.status_code == 422, (field, r.status_code, r.text)
    lead = await make_lead(admin, "Анна")
    r = await admin.patch(f"{L}/{lead['id']}", json={"comment": "x\x00"})
    assert r.status_code == 422
    r = await admin.post(CA, json={"name": "a\x00", "branch_id": branch,
                                   "platform": "Google", "start_date": "2026-09-01"})  # fmt: skip
    assert r.status_code == 422
    assert (await sql("SELECT count(*) FROM lead"))[0][0] == 1
    assert (await admin.get(L, params={"q": "a\x00"})).status_code == 422


async def test_convert_needs_students_change(admin: httpx.AsyncClient, login_as: SignedIn) -> None:
    """Converting creates or links a student: a leads-only role must not create students or
    learn which student owns a phone (the duplicate warning hides students from it too)."""
    await make_student("Скрытый Студент", phone="901112233")
    lead = await make_lead(admin, "Анна", phone="901112233")
    caller = await login_as("caller", await make_role("Колл-центр", {"leads": "change"}))
    r = await caller.post(f"{L}/{lead['id']}/convert", json={})
    assert r.status_code == 403
    assert "Скрытый" not in r.text
    assert (await sql("SELECT status FROM lead"))[0][0] != "записан"
    viewer = await login_as(
        "sv", await make_role("Лиды+студенты чтение", {"leads": "change", "students": "view"})
    )
    assert (await viewer.post(f"{L}/{lead['id']}/convert", json={})).status_code == 403


async def test_import_row_limit_stops_parsing_early() -> None:
    """The row cap is enforced while the CSV is read, not after the whole table is built."""
    from app.core.errors import DomainError
    from app.leads import importer

    rows = [["full_name"], *([["a"]] * (importer.MAX_ROWS + 5))]
    data = csv_bytes(rows) + b'\n"' + b"x" * 200_000  # a field past the csv size limit at the end
    try:
        importer.parse_csv(data)
    except DomainError as e:
        assert e.error.code == "import_too_many_rows", e.error.code
    else:
        raise AssertionError("no error")


async def test_import_bad_content_length_is_not_a_500(admin: httpx.AsyncClient) -> None:
    branch = str(await branch_id("Душанбе"))
    r = await admin.post(
        f"{L}/import/preview",
        content=csv_bytes([["full_name"], ["Анна"]]),
        params={"branch_id": branch},
        headers={"Content-Type": "text/csv", "Content-Length": "abc"},
    )
    assert r.status_code < 500


async def test_import_needs_csrf_header_and_stays_in_own_branch(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    from tests.sales import manager_role

    anna = await login_as("anna", await manager_role(), branch="Душанбе")
    data = csv_bytes([["full_name", "phone"], ["Анна", "901112233"]])
    r = await anna.post(
        f"{L}/import", content=data, headers={"Content-Type": "text/plain",
                                              "X-Requested-With": ""},
    )  # fmt: skip
    assert r.status_code == 403
    ist = str(await branch_id("Истаравшан"))
    r = await _csv(anna, "/import", data, branch_id=ist)
    assert r.status_code == 403
    foreign = await make_campaign(admin, branch="Истаравшан")
    r = await _csv(anna, "/import", data, campaign_id=foreign["id"])
    assert r.status_code == 404
    assert (await sql("SELECT count(*) FROM lead"))[0][0] == 0
    # preview with another branch's phones reveals nothing
    await make_lead(admin, "Чужой", phone="901112233", branch="Истаравшан")
    await make_student("Чужой студент", branch="Истаравшан", phone="901112233")
    r = await _csv(anna, "/import/preview", data)
    assert r.json()["rows"][0]["duplicates"] == []


async def test_sales_metrics_and_exports_are_branch_scoped(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    await make_campaign(admin, "Чужая", branch="Истаравшан", budget=5000)
    await make_lead(admin, "Чужой", branch="Истаравшан")
    role = await make_role("Продажи", {"leads": "view", "marketing": "view", "dashboard": "view"})
    anna = await login_as("anna2", role, branch="Душанбе")
    ist = str(await branch_id("Истаравшан"))
    assert (await anna.get(f"{M}/leads", params={"branch_id": ist})).json()["counts"]["всего"] == 0
    assert (await anna.get(f"{M}/finance")).json()["ads"] in ("0.00", 0, "0")
    assert (await anna.get(CA, params={"branch_id": ist})).json()["items"] == []
    assert (await anna.get(L, params={"branch_id": ist})).json()["total"] == 0
    for url in (f"{L}/export", f"{CA}/export"):
        r = await anna.get(url)
        assert r.status_code == 200
        ws = load_workbook(BytesIO(r.content)).active
        assert ws is not None and ws.max_row == 1  # header only
