"""Security of the import from the old system: the old data is untrusted input, the importer is a
CLI only, the report holds no personal data, imported accounts cannot sign in. Invented data only."""

import copy
import io
import json
import os
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from app.core import security
from app.core.config import settings
from app.core.db import sessionmaker
from app.importer import __main__ as cli
from app.importer import engine as importer_engine
from app.importer import parse as p
from app.importer import run_import
from app.main import app
from tests.conftest import Client, sql

FIXTURE = Path(__file__).parent / "fixtures" / "old_snapshot.json"
DATE = "Mon Feb 02 2026 12:00:00 GMT+0500 (Таджикистан)"
LONG = "Я" * 20_000
SECRET = "Секретов Персонаж"


@pytest.fixture
def snapshot() -> dict[str, list[dict[str, Any]]]:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def hostile(snapshot: dict[str, list[dict[str, Any]]]) -> dict[str, list[dict[str, Any]]]:
    """The fixture plus rows with every kind of garbage an old sheet can hold."""
    snap = copy.deepcopy(snapshot)
    name = '=HYPERLINK("x")\x00 <script>alert(1)</script> ‮Али​ Злодеев ' + LONG
    snap["STUDENTS"].append(
        {
            "student_id": "S-900",
            "full_name": name,
            "phone": "=1+1\x00" + LONG,
            "city": "Душанбе",
            "gender": "муж",
            "registration_date": DATE,
        }  # fmt: skip
    )
    snap["ENROLLMENTS"].append(
        {
            "enrollment_id": "E-900",
            "student_id": "S-900",
            "student_name": name,
            "cohort_id": "G-001",
            "status": "Активный",
            "final_score": "1e309",
        }  # fmt: skip
    )
    for i, (amount, month) in enumerate(
        [
            ("1e309", "1"),
            ("NaN", "1"),
            ("-5", "1"),
            ("99999999999999", "1"),
            ("Infinity", "1"),
            ("sNaN", "1"),
            ("1e-400", "1"),
            ("100", "1e309"),
            ("100", "Infinity"),
            ("100", "sNaN"),
            ("100", "99999"),
        ]  # fmt: skip
    ):
        snap["PAYMENTS"].append(
            {
                "payment_id": f"P-9{i:02}",
                "student_id": "S-900",
                "student_name": name,
                "cohort_id": "G-001",
                "payment_type": "Помесячно" if i else "=SUM(A1)\x00",
                "month_number": month,
                "amount": amount,
                "payment_method": "Alif",
                "payment_date": DATE,
                "verified": "Да",
                "notes": "@cmd\x00" + LONG,
                "screenshot_link": "+" + LONG,
            }  # fmt: skip
        )
    snap["GRADES"] += [
        {"grade_id": "GR-900", "student_id": "S-900", "student_name": name, "cohort_id": "G-001",
         "assessment_type": "homework", "score": "-5", "max_score": "20"},
        {"grade_id": "GR-901", "student_id": "S-900", "student_name": name, "cohort_id": "G-001",
         "assessment_type": "final_work", "score": "1e309", "max_score": "-60"},
    ]  # fmt: skip
    snap["LEADS"].append(
        {
            "lead_id": "L-900",
            "timestamp": "Mon Feb 02 2026 12:00:00 GMT+9900",
            "full_name": name,
            "phone": "=cmd|' /C calc'!A0" + LONG,
            "comments": "-2+3\x00" + LONG,
            "city": "Душанбе",
            "source": "Платный",
        }  # fmt: skip
    )
    snap["INSTRUCTORS"].append(
        {
            "instructor_id": "I-900",
            "full_name": "Злодей Инструкторов\x00" + LONG,
            "phone": "\x00",
            "email": "=evil@x.tj" + LONG,
            "city": "Душанбе",
            "salary_per_month": "1e309",
            "status": "Активный",
        }  # fmt: skip
    )
    snap["COURSES"].append(
        {
            "course_id": "C-900",
            "course_name": "Злой курс\x00" + LONG,
            "level": "",
            "price_per_month": "1e309",
            "duration_months": "Infinity",
            "city": "Душанбе",
        }  # fmt: skip
    )
    snap["MARKETING"].append(
        {
            "campaign_id": "MKT-900",
            "course_name": "=EVIL()" + LONG,
            "city": "Душанбе",
            "platform": "Facebook Ads",
            "start_date": DATE,
            "budget_tjs": "1e20",
        }  # fmt: skip
    )
    cohort = dict(snap["COHORTS"][0])
    cohort.update(cohort_id="G-900", room="=ROOM\x00" + LONG, cost="1e309", max_students="1e309")
    snap["COHORTS"].append(cohort)
    return snap


async def do_import(snapshot: dict[str, Any], *, apply: bool) -> Any:
    async with sessionmaker() as db:
        return await run_import(db, snapshot, apply=apply, source="old_snapshot.json")


# ---- untrusted values ---------------------------------------------------------------------------


def test_numbers_that_are_not_money_or_counts_are_none() -> None:
    for raw in ("1e309", "Infinity", "-Infinity", "NaN", "sNaN", "1e-400", "99999999999999"):
        assert p.to_decimal(raw) is None, raw
        assert p.to_int(raw) is None, raw
    assert p.to_decimal("600.500") == Decimal("600.5")
    assert p.to_decimal("9999999999.99") == Decimal("9999999999.99")
    assert p.to_int("1e6") is None and p.to_int("36") == 36


def test_a_garbage_utc_offset_is_not_a_date() -> None:
    assert p.parse_dt("Mon Feb 02 2026 12:00:00 GMT+9900") is None
    assert p.parse_date("Mon Feb 02 2026 12:00:00 GMT-2400") is None


def test_text_drops_nul_and_control_characters() -> None:
    assert p.text("Ал\x00и\x07 \x1bВали") == "Али Вали"
    assert p.text("строка\nдругая") == "строка\nдругая"


async def test_hostile_snapshot_imports_within_the_app_limits(
    snapshot: dict[str, Any], tmp_path: Path
) -> None:
    report = await do_import(hostile(snapshot), apply=True)
    limits = {
        ("student", "full_name"): 200, ("student", "phone"): 32, ("lead", "full_name"): 200,
        ("lead", "raw_name"): 300, ("lead", "phone"): 40, ("lead", "phone_raw"): 40,
        ("lead", "comment"): 2000, ("payment", "note"): 500, ("instructor_payout", "note"): 500,
        ("app_user", "full_name"): 200, ("app_user", "email"): 200, ("app_user", "login"): 64,
        ("room", "name"): 100, ("course", "name"): 200, ("campaign", "name"): 200,
        ("attendance", "reason"): 200,
    }  # fmt: skip
    for (table, column), limit in limits.items():
        longest = (await sql(f"SELECT coalesce(max(length({column})), 0) FROM {table}"))[0][0]
        assert longest <= limit, (table, column, longest)
    # no stored free text starts a spreadsheet formula
    for table, column in (("payment", "note"), ("lead", "comment"), ("campaign", "name"),
                          ("room", "name"), ("app_user", "email"), ("lead", "phone")):  # fmt: skip
        rows = await sql(f"SELECT {column} FROM {table} WHERE {column} ~ '^[=+@\\-]'")
        assert not [r for r in rows if not re.fullmatch(r"\+\d+", r[0])], (table, column, rows)
    amounts = [r[0] for r in await sql("SELECT amount FROM payment")]
    assert all(Decimal(0) < a < Decimal("10000000000") for a in amounts)
    months = [r[0] for r in await sql("SELECT month_no FROM payment")]
    assert all(1 <= m <= 36 for m in months)
    # the report files are written even with control characters in the review data
    md, xlsx = report.write(tmp_path, "h")
    wb = load_workbook(xlsx)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", c.value
    assert "\x00" not in md.read_text(encoding="utf-8")


async def test_exports_of_imported_rows_hold_no_formulas(
    snapshot: dict[str, Any], admin: Client
) -> None:
    await do_import(hostile(snapshot), apply=True)
    for url in ("/api/v1/leads/export", "/api/v1/students/export", "/api/v1/payments/export"):
        r = await admin.get(url)
        assert r.status_code == 200, (url, r.text)
        wb = load_workbook(io.BytesIO(r.content))
        for ws in wb:
            for row in ws.iter_rows():
                for c in row:
                    assert c.data_type != "f", (url, c.value)
                    assert not (isinstance(c.value, str) and c.value[:1] in "=@"), (url, c.value)


# ---- the report and the files -----------------------------------------------------------------


def _personal_values(snapshot: dict[str, list[dict[str, Any]]]) -> set[str]:
    keys = {"full_name", "student_name", "instructor_name", "phone", "email", "notes",
            "comments", "instructor_note", "instructor_notes", "drop_reason"}  # fmt: skip
    found: set[str] = set()
    for sheet, rows in snapshot.items():
        for r in rows:
            for k, v in r.items():
                # INSTRUCTOR_PAYMENTS columns are shifted: cohort_id holds the name, instructor_name an id
                shifted = sheet == "INSTRUCTOR_PAYMENTS"
                personal = (k in keys and not (shifted and k == "instructor_name")) or (
                    shifted and k == "cohort_id"
                )
                value = p.text(v)
                if personal and len(value) >= 4:
                    found |= {value, p.name_of(value)[0]}
                    digits = re.sub(r"\D", "", value)
                    if len(digits) >= 9:
                        found.add(digits[-9:])
    return found


async def test_report_md_holds_no_name_phone_or_email_of_the_snapshot(
    snapshot: dict[str, Any], tmp_path: Path
) -> None:
    report = await do_import(snapshot, apply=False)
    md = report.write(tmp_path, "t")[0].read_text(encoding="utf-8")
    leaked = sorted(v for v in _personal_values(snapshot) if v in md)
    assert not leaked, leaked


async def test_report_ids_cannot_inject_markdown_or_flood(snapshot: dict[str, Any]) -> None:
    snap = copy.deepcopy(snapshot)
    snap["LEADS"].append({"lead_id": "L|1\n# fake | `x` " + LONG, "full_name": ""})
    report = await do_import(snap, apply=False)
    md = report.markdown()
    line = next(x for x in md.splitlines() if "LEADS** — нет ФИО" in x)
    assert len(line) < 300 and "`x`" not in line and "\n# fake" not in md


async def test_report_files_are_private_and_default_to_import_data(
    snapshot: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(cli, "DEFAULT_OUT_DIR", tmp_path / "import-data")
    old = os.umask(0o022)
    try:
        assert await cli.main([str(FIXTURE), "--dry-run"]) == 0
    finally:
        os.umask(old)
    written = list((tmp_path / "import-data").iterdir())
    assert len(written) == 2
    for f in written:
        assert f.stat().st_mode & 0o777 == 0o600, f
    assert (tmp_path / "import-data").stat().st_mode & 0o077 == 0
    assert cli.DEFAULT_OUT_DIR.name == "import-data"  # gitignored at the repo root


async def test_a_failure_prints_no_row_content(
    snapshot: dict[str, Any],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    async def boom(self: Any) -> None:
        raise ValueError(f"bad row: {SECRET} 931234567")

    monkeypatch.setattr(importer_engine.Importer, "_leads", boom)
    rc = await cli.main([str(FIXTURE), "--dry-run", "--out-dir", str(tmp_path)])
    out = capsys.readouterr()
    assert rc != 0
    assert SECRET not in out.out + out.err and "931234567" not in out.out + out.err
    assert "ValueError" in out.err


async def test_stdout_lists_no_rows(
    snapshot: dict[str, Any], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert await cli.main([str(FIXTURE), "--dry-run", "--out-dir", str(tmp_path)]) == 0
    printed = capsys.readouterr()
    leaked = [v for v in _personal_values(snapshot) if v in printed.out + printed.err]
    assert not leaked, leaked


# ---- reachable only from the command line, guarded in production ------------------------------


def test_no_http_route_runs_the_import() -> None:
    paths = [getattr(r, "path", "") for r in app.routes]
    assert not [x for x in paths if "import" in x.lower() or "legacy" in x.lower()]


async def test_apply_in_production_needs_the_database_name_typed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "env", "production")
    db_name = settings.database_url.rsplit("/", 1)[1].split("?")[0]
    before = (await sql("SELECT count(*) FROM student"))[0][0]
    args = [str(FIXTURE), "--apply", "--out-dir", str(tmp_path)]
    assert await cli.main(args) == 2
    assert await cli.main([*args, "--confirm-db", "ilmhona-wrong"]) == 2
    assert (await sql("SELECT count(*) FROM student"))[0][0] == before
    assert await cli.main([*args, "--confirm-db", db_name]) == 0
    assert (await sql("SELECT count(*) FROM student"))[0][0] == 8
    # a dry-run never needs it
    assert await cli.main([str(FIXTURE), "--dry-run", "--out-dir", str(tmp_path)]) == 0


# ---- imported accounts -------------------------------------------------------------------------


async def test_imported_users_are_inactive_instructors_that_cannot_sign_in(
    snapshot: dict[str, Any], client: Client
) -> None:
    await do_import(snapshot, apply=True)
    rows = await sql(
        "SELECT u.login, u.is_active, u.password_hash, r.code FROM app_user u "
        "JOIN role r ON r.id = u.role_id WHERE u.legacy_ref IS NOT NULL"
    )
    assert rows
    assert {r[3] for r in rows} == {"instructor"}
    assert (await sql("SELECT count(*) FROM app_user u JOIN role r ON r.id = u.role_id "
                      "WHERE r.code = 'superadmin'"))[0][0] == 0  # fmt: skip
    unknown = await client.post("/api/v1/auth/login", json={"login": "nobody.x", "password": "x"})
    for login, active, pw_hash, _ in rows:
        assert active is False and not pw_hash.startswith("$argon2")
        for guess in ("", "x", pw_hash):
            r = await client.post("/api/v1/auth/login", json={"login": login, "password": guess})
            if guess:
                assert (r.status_code, r.json()) == (unknown.status_code, unknown.json())


def test_an_unusable_hash_costs_a_real_check(monkeypatch: pytest.MonkeyPatch) -> None:
    """An imported login (guessable from the name) must not answer faster than an unknown one."""
    calls: list[str] = []
    real = security._hasher

    class Counting:
        def verify(self, stored: str, password: str) -> bool:
            calls.append(stored)
            return real.verify(stored, password)

    monkeypatch.setattr(security, "_hasher", Counting())
    assert security.verify_password(importer_engine.UNUSABLE_HASH, "guess") is False
    assert security._DUMMY in calls


# ---- legacy_ref and the audit log --------------------------------------------------------------


async def test_legacy_ref_is_not_in_api_responses(snapshot: dict[str, Any], admin: Client) -> None:
    await do_import(snapshot, apply=True)
    sid = (await sql("SELECT id FROM student WHERE legacy_ref IS NOT NULL LIMIT 1"))[0][0]
    for url in ("/api/v1/students", f"/api/v1/students/{sid}", "/api/v1/leads", "/api/v1/users",
                "/api/v1/cohorts", "/api/v1/payments", "/api/v1/rooms", "/api/v1/courses",
                "/api/v1/campaigns", "/api/v1/audit"):  # fmt: skip
        r = await admin.get(url)
        assert r.status_code == 200, (url, r.text)
        assert "legacy" not in r.text and "STUDENTS:" not in r.text, url


async def test_one_audit_entry_per_apply_with_counts_only(snapshot: dict[str, Any]) -> None:
    await do_import(snapshot, apply=True)
    rows = await sql(
        "SELECT user_id, diff::text, action FROM audit_log WHERE action = 'import.apply'"
    )
    assert len(rows) == 1
    actor, diff, _ = rows[0]
    assert actor is None
    parsed = json.loads(diff)
    assert all(isinstance(v, list) and all(isinstance(n, int) for n in v) for v in parsed.values())
    # the cohort closing (the daily job's path) audits counts and ids too: no person anywhere
    every = " ".join(r[0] for r in await sql("SELECT diff::text FROM audit_log"))
    leaked = [v for v in _personal_values(snapshot) if v in every]
    assert not leaked, leaked
