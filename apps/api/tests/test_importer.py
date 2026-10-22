"""The import from the old system (docs/IMPORT.md) on an invented snapshot that has every quirk of
the real export: repeated ids for different people, JS date strings, the +11:25 lesson times, the
shifted INSTRUCTOR_PAYMENTS columns, `score:A|B|C` strings, Удалён and empty-id rows, room
variants and the same instructor in Latin and Cyrillic. No real data anywhere."""

import json
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from app.audit.sentences import Entry, sentence
from app.core.db import sessionmaker
from app.core.security import verify_password
from app.importer import engine as importer_engine
from app.importer import run_import
from app.importer.__main__ import main as importer_main
from app.importer.report import Report
from app.metrics.money import income_total
from tests.conftest import sql

FIXTURE = Path(__file__).parent / "fixtures" / "old_snapshot.json"
TABLES = (
    "room", "course", "app_user", "cohort", "lesson", "student", "enrollment", "payment",
    "instructor_payout", "campaign", "lead", "attendance", "grade", "audit_log",
)  # fmt: skip


@pytest.fixture
def snapshot() -> dict[str, list[dict[str, Any]]]:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


async def do_import(snapshot: dict[str, Any], *, apply: bool) -> Report:
    async with sessionmaker() as db:
        return await run_import(db, snapshot, apply=apply, source="old_snapshot.json")


async def counts() -> dict[str, int]:
    return {t: (await sql(f"SELECT count(*) FROM {t}"))[0][0] for t in TABLES}


def files(directory: Path, pattern: str) -> int:
    return len(list(directory.glob(pattern)))


async def one(query: str, **params: Any) -> Any:
    rows = await sql(query, **params)
    assert len(rows) == 1, rows
    return rows[0][0] if len(rows[0]) == 1 else tuple(rows[0])


async def ref_rows(table: str) -> set[str]:
    return {r[0] for r in await sql(f"SELECT legacy_ref FROM {table} WHERE legacy_ref IS NOT NULL")}


# ---- the transaction: dry-run, apply, re-apply, rollback ---------------------------------------


async def test_dry_run_writes_nothing_and_reports(snapshot: dict[str, Any]) -> None:
    before = await counts()
    report = await do_import(snapshot, apply=False)
    assert await counts() == before
    assert report.mode == "dry-run"
    assert report.imported["STUDENTS"] == 9 and report.imported["ENROLLMENTS"] == 10
    assert report.recon, "the reconciliation is part of the dry-run"


async def test_apply_imports_in_one_go_and_reapply_adds_nothing(snapshot: dict[str, Any]) -> None:
    first = await do_import(snapshot, apply=True)
    after_first = await counts()
    assert after_first["student"] == 9 and after_first["enrollment"] == 10
    assert first.imported["STUDENTS"] == 9 and first.existing["STUDENTS"] == 0
    again = await do_import(snapshot, apply=True)
    after_second = await counts()
    # only the audit trail grows: every imported table is untouched
    assert {t: n for t, n in after_second.items() if t != "audit_log"} == {
        t: n for t, n in after_first.items() if t != "audit_log"
    }
    assert after_second["audit_log"] == after_first["audit_log"] + 1
    assert again.imported["STUDENTS"] == 0 and again.existing["STUDENTS"] == 9
    assert again.imported["PAYMENTS"] == 0 and again.imported["COHORTS"] == 0
    assert again.imported["ATTENDANCE"] == 0 and again.existing["ATTENDANCE"] == 7  # the PK


async def test_apply_writes_an_audit_entry_dry_run_does_not(snapshot: dict[str, Any]) -> None:
    await do_import(snapshot, apply=False)
    assert (await sql("SELECT count(*) FROM audit_log WHERE action = 'import.apply'"))[0][0] == 0
    await do_import(snapshot, apply=True)
    action, diff = await one("SELECT action, diff FROM audit_log WHERE action = 'import.apply'")
    assert action == "import.apply" and diff["student"] == [0, 9]
    text = sentence(
        Entry(actor=None, action="import.apply", entity="import", label=None, diff=diff)
    )
    assert "перенос из старой системы" in text and "9 студентов" in text


async def test_error_midway_rolls_everything_back(
    snapshot: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    async def boom(self: Any) -> None:
        raise RuntimeError("disk on fire")

    monkeypatch.setattr(importer_engine.Importer, "_leads", boom)
    before = await counts()
    with pytest.raises(RuntimeError):
        await do_import(snapshot, apply=True)
    assert await counts() == before


async def test_rows_without_legacy_ref_are_never_touched(snapshot: dict[str, Any]) -> None:
    await sql(
        "INSERT INTO course (name, default_price_month, duration_months, format) "
        "VALUES ('Alpha', 111, 5, 'online')"
    )
    await sql(
        "INSERT INTO student (full_name, branch_id, phone) "
        "SELECT 'Рукописный Студент', id, '+992900000000' FROM branch LIMIT 1"
    )
    await sql(
        "INSERT INTO room (branch_id, name, capacity) "
        "SELECT id, 'Jobs', 7 FROM branch WHERE name = 'Душанбе'"
    )
    before = {
        "course": await sql("SELECT id, name, default_price_month, duration_months, format, version, legacy_ref FROM course"),
        "student": await sql("SELECT id, full_name, phone, version, legacy_ref FROM student"),
        "room": await sql("SELECT id, name, capacity, version, legacy_ref FROM room WHERE name = 'Jobs'"),
    }  # fmt: skip
    await do_import(snapshot, apply=True)
    for table, query in (
        ("course", "SELECT id, name, default_price_month, duration_months, format, version, legacy_ref FROM course WHERE name = 'Alpha'"),
        ("student", "SELECT id, full_name, phone, version, legacy_ref FROM student WHERE full_name = 'Рукописный Студент'"),
        ("room", "SELECT id, name, capacity, version, legacy_ref FROM room WHERE name = 'Jobs'"),
    ):  # fmt: skip
        assert await sql(query) == before[table], table
    assert (await sql("SELECT count(*) FROM course WHERE name = 'Alpha'"))[0][0] == 1
    assert (await sql("SELECT count(*) FROM room WHERE name = 'Jobs'"))[0][0] == 1
    # the hand-made course carries the imported cohorts of that name
    alpha = await one("SELECT id FROM course WHERE name = 'Alpha'")
    assert (await sql("SELECT count(*) FROM cohort WHERE course_id = :c", c=alpha))[0][0] >= 3


async def test_every_imported_row_has_a_unique_legacy_ref(snapshot: dict[str, Any]) -> None:
    await do_import(snapshot, apply=True)
    for table in ("room", "course", "app_user", "cohort", "lesson", "student", "enrollment",
                  "payment", "instructor_payout", "campaign", "lead"):  # fmt: skip
        total = (await sql(f"SELECT count(*) FROM {table}"))[0][0]
        refs = await ref_rows(table)
        assert total == len(refs), table


# ---- courses, instructors, cohorts, lessons ----------------------------------------------------


async def test_courses_merge_levels_and_cohort_courses_are_found_by_prefix(
    snapshot: dict[str, Any],
) -> None:
    report = await do_import(snapshot, apply=True)
    names = {r[0] for r in await sql("SELECT name FROM course")}
    assert names == {"Alpha", "Beta Course", "Gamma Course"}  # level is not part of the name
    assert report.merged["COURSES"] == 1
    alpha = await one("SELECT id FROM course WHERE name = 'Alpha'")
    # G-002 has no course id, only «Alpha 1 (Feb)»
    code_course = await sql("SELECT c.legacy_ref, c.course_id FROM cohort c ORDER BY c.legacy_ref")
    assert {ref: cid for ref, cid in code_course}["COHORTS:G-002"] == alpha
    assert await one("SELECT format FROM course WHERE name = 'Beta Course'") == "offline"
    # a cohort whose course is not in COURSES makes one (flagged for review)
    assert any(r.reason.startswith("курс создан") for r in report.review)


async def test_instructors_inactive_one_per_person_latin_equals_cyrillic(
    snapshot: dict[str, Any],
) -> None:
    report = await do_import(snapshot, apply=True)
    rows = await sql(
        "SELECT u.login, u.full_name, u.is_active, u.must_change_password, u.default_rate, "
        "u.password_hash, u.phone, r.name, b.name, u.legacy_ref "
        "FROM app_user u JOIN role r ON r.id = u.role_id LEFT JOIN branch b ON b.id = u.branch_id "
        "WHERE u.legacy_ref IS NOT NULL ORDER BY u.full_name"
    )
    assert [r[1] for r in rows] == [
        "Aziz Testov",
        "Bahor Fakeva",
        "Eldor Sample",
    ]  # Dilshod: Удалён
    aziz, bahor, eldor = rows
    assert aziz[0] == "aziz.testov" and bahor[0] == "bahor.fakeva"
    for r in rows:
        assert r[2] is False and r[3] is True and r[7] == "Инструктор"
        assert verify_password(r[5], "anything") is False  # no password until activation
    assert aziz[4] == Decimal("2000") and bahor[4] == 0 and eldor[4] == 0
    assert aziz[6] == "+992931234567" and bahor[6] is None  # the phone is taken: not recorded
    assert (aziz[8], bahor[8]) == ("Душанбе", "Истаравшан")
    assert report.merged["INSTRUCTORS"] == 2 and report.dropped_count("INSTRUCTORS") == 1
    reasons = {r.reason for r in report.review if r.sheet == "INSTRUCTORS"}
    assert any("ставка 1" in x for x in reasons) and any("телефон уже занят" in x for x in reasons)
    # the cohort of the Cyrillic spelling belongs to the same account
    ins = await sql(
        "SELECT DISTINCT instructor_id FROM cohort WHERE legacy_ref IN ('COHORTS:G-001', 'COHORTS:G-002')"
    )
    assert len(ins) == 1


async def test_login_never_collides_with_an_existing_account(snapshot: dict[str, Any]) -> None:
    await sql(
        "INSERT INTO app_user (login, full_name, role_id, password_hash) "
        "SELECT 'aziz.testov', 'Занято', id, 'x' FROM role WHERE code = 'manager'"
    )
    await do_import(snapshot, apply=True)
    logins = {r[0] for r in await sql("SELECT login FROM app_user WHERE legacy_ref IS NOT NULL")}
    assert "aziz.testov2" in logins and "aziz.testov" not in logins


async def test_times_are_calibrated_and_bad_ones_leave_the_cohort_without_schedule(
    snapshot: dict[str, Any],
) -> None:
    report = await do_import(snapshot, apply=True)
    # 07:05 -> 18:30 and 08:35 -> 20:00 (G-001), 05:35 -> 17:00 and 07:05 -> 18:30 (G-003)
    t1 = await one("SELECT start_time, end_time FROM cohort WHERE legacy_ref = 'COHORTS:G-001'")
    t3 = await one("SELECT start_time, end_time FROM cohort WHERE legacy_ref = 'COHORTS:G-003'")
    assert (t1[0].strftime("%H:%M"), t1[1].strftime("%H:%M")) == ("18:30", "20:00")
    assert (t3[0].strftime("%H:%M"), t3[1].strftime("%H:%M")) == ("17:00", "18:30")
    # lessons are in local time and inside 07:00-22:00 (the constraint is the guarantee, #6)
    first = await one(
        "SELECT to_char(starts_at AT TIME ZONE 'Asia/Dushanbe', 'YYYY-MM-DD HH24:MI') FROM lesson l "
        "JOIN cohort c ON c.id = l.cohort_id WHERE c.legacy_ref = 'COHORTS:G-001' "
        "ORDER BY starts_at LIMIT 1"
    )
    assert first == "2026-02-02 18:30"
    # G-004 has `Sat Dec 30 1899` times: no schedule, no lessons, a review row
    assert (await sql("SELECT count(*) FROM lesson l JOIN cohort c ON c.id = l.cohort_id "
                      "WHERE c.legacy_ref = 'COHORTS:G-004'"))[0][0] == 0  # fmt: skip
    assert any(r.old_id == "G-004" and "расписани" in r.reason or "время уроков" in r.reason
               for r in report.review)  # fmt: skip


async def test_lessons_follow_the_schedule(snapshot: dict[str, Any]) -> None:
    await do_import(snapshot, apply=True)
    # G-001: Mon/Wed/Fri from 2026-02-02 to 2026-03-01 -> 12 generated lessons (+1 manual, below)
    days = await sql(
        "SELECT l.lesson_date, l.source FROM lesson l JOIN cohort c ON c.id = l.cohort_id "
        "WHERE c.legacy_ref = 'COHORTS:G-001' AND l.source = 'generated' ORDER BY 1"
    )
    assert len(days) == 12
    assert {d[0].isoweekday() for d in days} == {1, 3, 5}
    # G-003 (Вт/Чт/Сб): 9 weeks minus the start: Tuesdays, Thursdays, Saturdays
    assert {
        r[0].isoweekday()
        for r in await sql(
            "SELECT l.lesson_date FROM lesson l JOIN cohort c ON c.id = l.cohort_id "
            "WHERE c.legacy_ref = 'COHORTS:G-003'"
        )
    } == {2, 4, 6}


async def test_rooms_variants_conflicts_and_missing_rooms(snapshot: dict[str, Any]) -> None:
    report = await do_import(snapshot, apply=True)
    rooms = await sql(
        "SELECT c.legacy_ref, r.name, b.name FROM cohort c JOIN room r ON r.id = c.room_id "
        "JOIN branch b ON b.id = r.branch_id"
    )
    by_cohort = {ref: (room, branch) for ref, room, branch in rooms}
    assert by_cohort["COHORTS:G-001"] == ("Jobs", "Душанбе")
    # Istaravshan-1 is written the Cyrillic way and the cohort without a city lands in Истаравшан
    assert by_cohort["COHORTS:G-003"] == ("Истаравшан-1", "Истаравшан")
    # «Работа», an empty room and a room taken by G-001 at the same time: their own stand-in room
    for old in ("G-004", "G-006", "G-002"):
        assert by_cohort[f"COHORTS:{old}"][0].startswith("Без аудитории"), old
    reasons = " ".join(r.reason for r in report.review if r.sheet == "COHORTS")
    assert "«Работа»" in reasons and "аудитория занята" in reasons
    # #7 holds in the data: no two live lessons overlap in a room
    assert (await sql(
        "SELECT count(*) FROM lesson a JOIN lesson b ON a.room_id = b.room_id AND a.id < b.id "
        "AND tstzrange(a.starts_at, a.ends_at) && tstzrange(b.starts_at, b.ends_at) "
        "AND NOT a.cancelled AND NOT b.cancelled"
    ))[0][0] == 0  # fmt: skip


async def test_cohort_rows_dropped_merged_and_statuses(snapshot: dict[str, Any]) -> None:
    report = await do_import(snapshot, apply=True)
    statuses = {r[0]: r[1] for r in await sql("SELECT legacy_ref, status FROM cohort")}
    assert statuses == {
        "COHORTS:G-001": "завершён",
        "COHORTS:G-002": "завершён",
        "COHORTS:G-003": "завершён",
        "COHORTS:G-004": "завершён",
        "COHORTS:G-005": "отменён",
        "COHORTS:G-006": "набор",
    }  # G-007 Удалён, G-008 no start date, the empty-id row: not transferred
    reasons = {r for sheet in report.dropped.values() for r in sheet}
    assert {"нет ID", "статус Удалён", "нет даты начала"} <= reasons
    assert report.merged["COHORTS"] == 1  # the second G-006 row
    # the 15 students of the cohort: capacity from max_students, price from cost or the course
    cap, price = await one(
        "SELECT capacity, price_month FROM cohort WHERE legacy_ref = 'COHORTS:G-002'"
    )
    assert (cap, price) == (15, Decimal("600"))  # empty cost -> the course's price
    assert await one(
        "SELECT price_month FROM cohort WHERE legacy_ref = 'COHORTS:G-001'"
    ) == Decimal("650")
    # enrolled_count is ignored (#3): it is computed from enrollments


# ---- students and enrollments ------------------------------------------------------------------


async def test_students_merge_by_phone_or_name_and_an_id_is_not_a_person(
    snapshot: dict[str, Any],
) -> None:
    report = await do_import(snapshot, apply=True)
    names = sorted(r[0] for r in await sql("SELECT full_name FROM student"))
    assert names == sorted(
        ["Тест Студентов", "Ali Fakerov", "Мура Выдуманная", "Нора Вымышленная",
         "Рустам Несуществующий", "Зарина Придуманная", "Фейк Эмодзи", "Без Идентификатора", "Совсем Другой"]
    )  # fmt: skip
    # S-001, the repeated S-001 row and S-002 (same phone, name reordered): one person; S-010 has the same phone and
    # another name: a separate student, flagged
    assert report.merged["STUDENTS"] == 3  # repeated S-001, S-002 (phone), S-004 (name)
    reasons = " ".join(r.reason for r in report.review if r.sheet == "STUDENTS")
    assert "проверить слияние: то же ФИО без телефона" in reasons
    assert "общий телефон — проверить" in reasons
    assert "заполнить телефон" in reasons and "телефон не распознан" in reasons
    # S-005 was two different people: both transferred, each with their own enrollment
    assert (
        await sql(
            "SELECT count(*) FROM student WHERE full_name IN ('Мура Выдуманная', 'Нора Вымышленная')"
        )
    )[0][0] == 2
    # the emoji is gone from the name; created_at is the registration date (earliest of the rows)
    assert (
        await one(
            "SELECT to_char(created_at AT TIME ZONE 'Asia/Dushanbe', 'YYYY-MM-DD') FROM student WHERE full_name = 'Тест Студентов'"
        )
    ) == "2026-01-20"
    assert await one("SELECT gender FROM student WHERE full_name = 'Тест Студентов'") == "male"
    assert (
        await one("SELECT phone FROM student WHERE full_name = 'Тест Студентов'") == "+992931111111"
    )
    assert (
        await one("SELECT phone FROM student WHERE full_name = 'Рустам Несуществующий'") is None
    )  # #ERROR!


async def test_students_get_a_branch(snapshot: dict[str, Any]) -> None:
    await do_import(snapshot, apply=True)
    branches = {r[0]: r[1] for r in await sql(
        "SELECT s.full_name, b.name FROM student s JOIN branch b ON b.id = s.branch_id")}  # fmt: skip
    assert branches["Тест Студентов"] == "Душанбе"
    assert branches["Нора Вымышленная"] == "Истаравшан"
    # no city at all: the branch of the cohort they study in
    assert branches["Зарина Придуманная"] == "Истаравшан"


async def test_enrollments_merge_pick_senior_status_and_close_with_the_cohort(
    snapshot: dict[str, Any],
) -> None:
    report = await do_import(snapshot, apply=True)
    status = {
        (r[0], r[1]): r[2]
        for r in await sql(
            "SELECT s.full_name, c.legacy_ref, e.status FROM enrollment e "
            "JOIN student s ON s.id = e.student_id JOIN cohort c ON c.id = e.cohort_id"
        )
    }
    assert len(status) == 10 and report.merged["ENROLLMENTS"] == 1
    # E-001: Активный + Отчислен -> the senior one; total 85 on a finished cohort -> выпускник
    assert status[("Тест Студентов", "COHORTS:G-001")] == "выпускник"
    assert status[("Ali Fakerov", "COHORTS:G-001")] == "отчислен"
    # score:20|20|60 -> homework 40, final 40: with attendance that passes 70
    assert status[("Мура Выдуманная", "COHORTS:G-001")] == "выпускник"
    # total only (40), «Ожидает», a low score:10|10|30: the cohort finished, nobody reached 70
    assert status[("Нора Вымышленная", "COHORTS:G-003")] == "завершил"
    assert status[("Рустам Несуществующий", "COHORTS:G-001")] == "завершил"
    assert status[("Зарина Придуманная", "COHORTS:G-003")] == "завершил"
    # a cancelled cohort keeps its enrollments as they were
    assert status[("Тест Студентов", "COHORTS:G-005")] == "активный"
    reasons = {r for sheet in report.dropped.get("ENROLLMENTS", {}) for r in [sheet]}
    assert "пустой ID" in reasons and "поток удалён или не перенесён" in reasons
    assert "студент не найден в STUDENTS" in reasons
    # the earliest enrollment date of the merged rows
    assert (
        await one(
            "SELECT to_char(e.enrolled_at AT TIME ZONE 'Asia/Dushanbe', 'YYYY-MM-DD') FROM enrollment e "
            "JOIN student s ON s.id = e.student_id JOIN cohort c ON c.id = e.cohort_id "
            "WHERE s.full_name = 'Тест Студентов' AND c.legacy_ref = 'COHORTS:G-001'"
        )
        == "2026-01-30"
    )


async def test_grades_from_score_strings_and_the_grades_sheet(snapshot: dict[str, Any]) -> None:
    report = await do_import(snapshot, apply=True)
    grades = {
        r[0]: (r[1], r[2])
        for r in await sql(
            "SELECT s.full_name, g.homework, g.final FROM grade g "
            "JOIN enrollment e ON e.id = g.enrollment_id JOIN student s ON s.id = e.student_id"
        )
    }
    assert grades["Мура Выдуманная"] == (40, 40)  # score:20|20|60 on the old 20 / 20 / 60 scale
    assert grades["Зарина Придуманная"] == (20, 20)  # score:10|10|30, unnamed column, halves
    assert grades["Ali Fakerov"] == (40, 20)  # GRADES sheet: 20/20 homework, 30/60 final
    assert "Тест Студентов" not in grades  # only a total (85): homework and final stay empty
    assert report.dropped["GRADES"]["зачисление не найдено или нет оценки"] == ["GR-003"]
    assert report.notes["зачисление только с итоговым баллом (Д/З и финал пусты)"] >= 2


# ---- payments, attendance ----------------------------------------------------------------------


async def test_payments_dedupe_months_status_and_no_orphans(snapshot: dict[str, Any]) -> None:
    report = await do_import(snapshot, apply=True)
    rows = await sql(
        "SELECT p.legacy_ref, p.month_no, p.amount, p.method, p.status, p.note, "
        "p.confirmed_at IS NOT NULL FROM payment p ORDER BY p.legacy_ref"
    )
    assert len(rows) == 5 and report.merged["PAYMENTS"] == 1  # the exact repeat of P-001
    by_ref = {r[0]: r for r in rows}
    p1 = [r for r in rows if r[0].startswith("PAYMENTS:P-001")]
    assert sorted((r[1], r[2]) for r in p1) == [(1, Decimal("600")), (2, Decimal("650"))]
    p2 = by_ref[next(k for k in by_ref if k.startswith("PAYMENTS:P-002"))]
    assert (p2[1], p2[2], p2[3]) == (1, Decimal("1800"), "наличные")
    assert p2[5] == "За весь курс; 1+1"  # the type goes to the note; a formula is neutralised
    p3 = by_ref[next(k for k in by_ref if k.startswith("PAYMENTS:P-003"))]
    assert (p3[3], p3[4], p3[6]) == ("payvand", "ожидает", False)  # verified is empty
    dropped = report.dropped["PAYMENTS"]
    assert dropped["студент не найден"] == ["P-004"] and dropped["нет суммы"] == ["P-005"]
    async with sessionmaker() as db:
        from app.models import Cohort

        total = await income_total(db, [Cohort.legacy_ref.is_not(None)])
    assert total == Decimal("3350.00")  # 600 + 650 + 1800 + 300, the pending one does not count


async def test_attendance_creates_missing_lessons_and_resolves_conflicts(
    snapshot: dict[str, Any],
) -> None:
    report = await do_import(snapshot, apply=True)

    async def mark(student: str, day: str) -> tuple[bool, str | None]:
        return await one(
            "SELECT a.present, a.reason FROM attendance a JOIN lesson l ON l.id = a.lesson_id "
            "JOIN enrollment e ON e.id = a.enrollment_id JOIN student s ON s.id = e.student_id "
            "WHERE s.full_name = :n AND l.lesson_date = :d",
            n=student,
            d=__import__("datetime").date.fromisoformat(day),
        )

    assert await mark("Тест Студентов", "2026-02-02") == (True, None)  # present + note beats absent
    assert await mark("Тест Студентов", "2026-02-04") == (False, None)  # no note: the last row
    assert await mark("Ali Fakerov", "2026-02-02") == (False, "Болен")  # expelled: history stays
    # 2026-02-03 is a Tuesday: no such lesson in the schedule, a manual one is created
    day, source = await one(
        "SELECT l.lesson_date, l.source FROM lesson l JOIN cohort c ON c.id = l.cohort_id "
        "WHERE c.legacy_ref = 'COHORTS:G-001' AND l.lesson_date = DATE '2026-02-03'"
    )
    assert source == "manual"
    assert await mark("Тест Студентов", "2026-02-03") == (True, None)
    assert report.imported["ATTENDANCE"] == 7
    assert report.notes["конфликт отметок (присутствовал/отсутствовал на один урок)"] == 2
    dropped = report.dropped["ATTENDANCE"]
    assert dropped["статус пустой"] == ["AT-3"] and dropped["зачисление не найдено"] == ["AT-9"]
    assert report.merged["ATTENDANCE"] == 2


async def test_attendance_feeds_the_attendance_score(snapshot: dict[str, Any]) -> None:
    from app.metrics import attendance_score

    await do_import(snapshot, apply=True)
    present, marked = await one(
        "SELECT count(*) FILTER (WHERE a.present), count(*) FROM attendance a "
        "JOIN enrollment e ON e.id = a.enrollment_id JOIN student s ON s.id = e.student_id "
        "WHERE s.full_name = 'Мура Выдуманная'"
    )
    assert (present, marked) == (3, 3)
    assert attendance_score(3, 13) == 5  # computed, never stored (#12)


# ---- sales, money ------------------------------------------------------------------------------


async def test_campaigns_skip_tests_and_recompute_metrics(snapshot: dict[str, Any]) -> None:
    report = await do_import(snapshot, apply=True)
    rows = await sql(
        "SELECT c.name, c.platform, c.budget, c.status, c.end_date, b.name, co.name "
        "FROM campaign c JOIN branch b ON b.id = c.branch_id LEFT JOIN course co ON co.id = c.course_id "
        "ORDER BY c.start_date"
    )
    assert [r[0] for r in rows] == ["Alpha 0 (Feb)", "Beta Course 0 (Mar)"]  # «тест» is not moved
    assert rows[0][1:4] == ("Facebook Ads", Decimal("1200.00"), "завершена")
    assert rows[1][2:5] == (Decimal("350.50"), "активна", None)
    assert (rows[0][5], rows[1][5]) == ("Душанбе", "Истаравшан")
    assert (rows[0][6], rows[1][6]) == ("Alpha", "Beta Course")
    assert report.dropped["MARKETING"]["тестовая кампания"] == ["MKT-002"]
    # dates instead of numbers in cost_per_lead / cost_per_student: not moved, computed anew
    assert (
        report.notes["кампания с повреждёнными метриками (даты вместо чисел): пересчитаны заново"]
        == 1
    )
    assert (await sql("SELECT count(*) FROM lead WHERE campaign_id IS NOT NULL"))[0][0] == 0


async def test_leads_clean_map_and_merge(snapshot: dict[str, Any]) -> None:
    report = await do_import(snapshot, apply=True)
    leads = {
        r[0]: r
        for r in await sql(
            "SELECT full_name, raw_name, phone, phone_raw, source, call_status, status, comment, "
            "needs_review, branch_id IS NOT NULL, co.name FROM lead l LEFT JOIN course co ON co.id = l.course_id"
        )
    }
    assert len(leads) == 4 and report.merged["LEADS"] == 1  # same phone and day
    fake = leads["Лид Фейков"]
    assert fake[1] == "💥 Лид Фейков 💥" and fake[2] == "+992935555555"  # the raw name stays
    assert (fake[4], fake[5], fake[6], fake[10]) == ("Facebook", "не звонили", "новый", "Alpha")
    assert fake[8] is True  # no city: to review
    ivan = leads["Иван Вымышленный"]
    assert (ivan[4], ivan[5], ivan[6]) == ("другое", "не ответил", "записан")  # Оплатил
    assert ivan[7].startswith("HYPERLINK") and "электронная почта" in ivan[7]  # no formula, a note
    assert leads["Без Телефонова"][2] is None and leads["Без Телефонова"][8] is True
    short = leads["Короткий Номер"]
    assert short[2] == "1234" and short[3] == "1234" and short[8] is True


async def test_instructor_payouts_are_read_with_the_shifted_columns(
    snapshot: dict[str, Any],
) -> None:
    report = await do_import(snapshot, apply=True)
    rows = await sql(
        "SELECT p.legacy_ref, u.full_name, c.legacy_ref, p.period, p.amount, p.paid_at, p.method "
        "FROM instructor_payout p JOIN app_user u ON u.id = p.instructor_id "
        "JOIN cohort c ON c.id = p.cohort_id ORDER BY c.legacy_ref"
    )
    assert [(r[1], r[2], r[4]) for r in rows] == [
        ("Aziz Testov", "COHORTS:G-001", Decimal("2000")),
        ("Aziz Testov", "COHORTS:G-002", Decimal("2000")),  # the Cyrillic spelling, same person
        ("Bahor Fakeva", "COHORTS:G-003", Decimal("800")),
    ]
    first = rows[0]
    assert str(first[3]) == "2026-02-01" and str(first[5]) == "2026-02-28"  # month 1 of the cohort
    assert report.merged["INSTRUCTOR_PAYMENTS"] == 1  # the exact repeat of IP-001
    dropped = report.dropped["INSTRUCTOR_PAYMENTS"]
    assert dropped["инструктор не совпал с инструктором потока"] == ["IP-004"]
    assert dropped["поток не найден"] == ["IP-005"] and dropped["пустой ID"] == ["—"]
    # profit reads the payouts (#9): the same code path as everywhere
    async with sessionmaker() as db:
        from app.metrics.money import payroll_total
        from app.models import Cohort

        assert await payroll_total(db, [Cohort.legacy_ref.is_not(None)]) == Decimal("4800.00")


# ---- the report --------------------------------------------------------------------------------


async def test_report_has_counts_and_ids_only_and_review_has_the_data(
    snapshot: dict[str, Any], tmp_path: Path
) -> None:
    report = await do_import(snapshot, apply=False)
    md_path, xlsx_path = report.write(tmp_path, "t")
    md = md_path.read_text(encoding="utf-8")
    for private in (
        "Тест Студентов",
        "Мура",
        "Aziz",
        "Азиз",
        "931111111",
        "935555555",
        "Ali Fakerov",
    ):
        assert private not in md, private
    assert "| STUDENTS | 13 | 9 |" in md  # read / imported
    assert "Сверка с интерфейсом старой системы" in md and "783 / 931 / 965" in md
    assert "390 / 400 / 451" in md
    wb = load_workbook(xlsx_path)
    assert "STUDENTS" in wb.sheetnames
    flat = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row)
    assert "Ali Fakerov" in flat  # the review list holds the people: it stays in import-data/
    assert md_path.stat().st_mode & 0o777 == 0o600


async def test_reconciliation_numbers_come_from_the_metrics(snapshot: dict[str, Any]) -> None:
    report = await do_import(snapshot, apply=False)
    recon = {r.label: r for r in report.recon}
    assert recon["Студентов (всего)"].new == "9"
    assert recon["Выпускников"].new == "2"
    assert recon["Сумма подтверждённых оплат, TJS"].new == "3350.00"


async def test_cli_dry_run_and_apply(snapshot: dict[str, Any], tmp_path: Path) -> None:
    before = await counts()
    assert await importer_main([str(FIXTURE), "--dry-run", "--out-dir", str(tmp_path)]) == 0
    assert await counts() == before
    assert files(tmp_path, "report-*.md") == 1 and files(tmp_path, "review-*.xlsx") == 1
    assert await importer_main([str(FIXTURE), "--apply", "--out-dir", str(tmp_path)]) == 0
    assert (await sql("SELECT count(*) FROM student"))[0][0] == 9
    with pytest.raises(SystemExit):
        await importer_main([str(FIXTURE), "--out-dir", str(tmp_path)])  # a mode is required


async def test_payment_without_enrollment_restores_it_and_the_money_balances(
    snapshot: dict[str, Any],
) -> None:
    report = await do_import(snapshot, apply=True)
    # P-006: the student and the cohort exist, the enrollment id does not -> made from the payment
    status, manual = await one(
        "SELECT e.status, e.status_manual FROM enrollment e JOIN student s ON s.id = e.student_id "
        "JOIN cohort c ON c.id = e.cohort_id WHERE s.full_name = 'Зарина Придуманная' "
        "AND c.legacy_ref = 'COHORTS:G-001'"
    )
    assert (status, manual) == ("завершил", True)
    assert any(r.reason == "зачисление восстановлено из оплаты" for r in report.review)
    # a payment of a cohort that is not imported is dropped with its amount in the review list
    # the ledger: sheet total - drops - exact repeats = imported, to the dirham
    (title, lines), *_ = report.ledgers
    assert title == "Сверка денег"
    by = {label: (n, amount) for label, n, amount in lines}
    assert by["Лист PAYMENTS: все строки / подтверждённая сумма"] == ("8", "4550.00")
    assert by["= перенесено (расчёт)"][1] == "3350.00"
    assert by["разница (должна быть 0)"][1] == "0.00"
    assert by["− слито: точные повторы строки (все поля, вкл. ID)"] == ("1", "600.00")


# ---- one phone: the same person or siblings ----------------------------------------------------


def _student(sid: str, name: str, phone: str) -> dict[str, Any]:
    return {
        "student_id": sid, "full_name": name, "phone": phone, "city": "Душанбе",
        "registration_date": "Tue Jan 20 2026 12:00:00 GMT+0500 (Таджикистан)", "status": "Активный",
    }  # fmt: skip


async def test_shared_phone_siblings_stay_apart_and_spellings_of_one_person_merge(
    snapshot: dict[str, Any],
) -> None:
    base = await do_import(snapshot, apply=False)
    snapshot["STUDENTS"] += [
        _student("S-101", "Ali Karimov", "937770001"),
        _student("S-102", "Алӣ Каримов", "937770001"),  # Latin = Cyrillic
        _student("S-103", "Каримов Али", "937770001"),  # word order
        _student("S-104", "Ali Karimof", "937770001"),  # one letter off
        _student("S-105", "Aziza Karimov", "937770001"),  # sister: same phone, same surname
        _student("S-106", "Зафар Рахимов", "937770002"),
        _student("S-107", "Зарина Султонова", "937770002"),
    ]
    for sid, name, eid, pid, amount in (
        ("S-105", "Aziza Karimov", "E-901", "P-901", "200"),
        ("S-106", "Зафар Рахимов", "E-902", "P-902", "300"),
        ("S-107", "Зарина Султонова", "E-903", "P-903", "400"),
    ):
        snapshot["ENROLLMENTS"].append(
            {"enrollment_id": eid, "student_id": sid, "student_name": name, "cohort_id": "G-001",
             "enrollment_date": "Fri Jan 30 2026 12:00:00 GMT+0500 (Таджикистан)", "status": "Активный"}
        )  # fmt: skip
        snapshot["PAYMENTS"].append(
            {"payment_id": pid, "enrollment_id": eid, "student_id": sid, "student_name": name,
             "cohort_id": "G-001", "payment_type": "Помесячно", "month_number": "1",
             "amount": amount, "payment_method": "Наличные", "verified": "Да",
             "payment_date": "Mon Feb 02 2026 12:00:00 GMT+0500 (Таджикистан)"}
        )  # fmt: skip
        snapshot["ATTENDANCE"].append(
            {"attendance_id": f"AT-{eid}", "enrollment_id": eid, "student_id": sid,
             "student_name": name, "cohort_id": "G-001", "lesson_date":
             "Mon Feb 02 2026 12:00:00 GMT+0500 (Таджикистан)", "lesson_number": "1",
             "status": "Присутствовал"}
        )  # fmt: skip
    report = await do_import(snapshot, apply=True)
    # Ali (4 spellings) + Aziza + Zafar + Zarina = 4 new students; 3 spellings merged into Ali
    assert report.imported["STUDENTS"] == base.imported["STUDENTS"] + 4
    assert report.merged["STUDENTS"] == base.merged["STUDENTS"] + 3
    reasons = dict(await sql("SELECT full_name, review_reason FROM student"))
    assert reasons["Ali Karimov"] == reasons["Aziza Karimov"] == "общий телефон — проверить"
    assert reasons["Зафар Рахимов"] == reasons["Зарина Султонова"] == "общий телефон — проверить"
    for name in ("Ali Karimov", "Aziza Karimov", "Зафар Рахимов", "Зарина Султонова"):
        assert await one("SELECT count(*) FROM student WHERE full_name = :n", n=name) == 1
    assert await one("SELECT count(*) FROM student WHERE phone = '+992937770001'") == 2
    # each child has exactly their own enrollment, payment and attendance
    for name, amount in (("Aziza Karimov", 200), ("Зафар Рахимов", 300), ("Зарина Султонова", 400)):
        assert await one(
            "SELECT count(*), sum(p.amount) FROM enrollment e JOIN student s ON s.id = e.student_id "
            "JOIN payment p ON p.enrollment_id = e.id WHERE s.full_name = :n",
            n=name,
        ) == (1, Decimal(amount))
        assert (
            await one(
                "SELECT count(*) FROM attendance a JOIN enrollment e ON e.id = a.enrollment_id "
                "JOIN student s ON s.id = e.student_id WHERE s.full_name = :n",
                n=name,
            )
            == 1
        )
    (_, lines), *_ = report.ledgers
    by = {label: amount for label, _n, amount in lines}
    assert by["разница (должна быть 0)"] == "0.00"
    assert by["= перенесено (расчёт)"] == "4250.00"  # 3350 + 200 + 300 + 400


async def test_instructor_overlaps_are_marked_not_fatal(snapshot: dict[str, Any]) -> None:
    """Two cohorts of one instructor at the same time: both sides marked «конфликт инструктора»
    and flagged; the import does not fail on lesson_instructor_no_overlap."""
    clash = {**next(r for r in snapshot["COHORTS"] if r["cohort_id"] == "G-001")}
    clash.update(cohort_id="G-901", room="", _rowIndex=901)
    snapshot["COHORTS"].append(clash)
    report = await do_import(snapshot, apply=True)
    marked = await sql(
        "SELECT DISTINCT c.legacy_ref FROM lesson l JOIN cohort c ON c.id = l.cohort_id "
        "WHERE l.instructor_conflict"
    )
    assert {r[0] for r in marked} >= {"COHORTS:G-001", "COHORTS:G-901"}
    flags = [r for r in report.review if "конфликт инструктора" in r.reason]
    assert "G-901" in {r.old_id for r in flags}
    assert (await sql(
        "SELECT count(*) FROM lesson a JOIN lesson b ON a.instructor_id = b.instructor_id "
        "AND a.id < b.id AND tstzrange(a.starts_at, a.ends_at) && tstzrange(b.starts_at, b.ends_at) "
        "AND NOT a.cancelled AND NOT b.cancelled "
        "AND NOT (a.instructor_conflict AND b.instructor_conflict)"
    ))[0][0] == 0  # fmt: skip
