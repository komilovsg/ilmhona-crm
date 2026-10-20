"""The табель: the attendance grid as an Excel file (same numbers as the screen)."""

from io import BytesIO

import httpx
from openpyxl import load_workbook

from tests.conftest import SignedIn, sql
from tests.edu import C
from tests.test_attendance import enroll, mark, past_cohort


async def test_tabel_export_equals_the_grid(admin: httpx.AsyncClient, login_as: SignedIn) -> None:
    cohort, lessons = await past_cohort(admin)
    a, b = await enroll(admin, cohort, "Алиев"), await enroll(admin, cohort, "Бобоев")
    await mark(admin, lessons[0], a)
    await mark(admin, lessons[1], a, present=False, reason="болел")
    await mark(admin, lessons[0], b)
    await admin.post(f"/api/v1/lessons/{lessons[2]}/cancel")

    grid = (await admin.get(f"{C}/{cohort['id']}/attendance")).json()
    r = await admin.get(f"{C}/{cohort['id']}/attendance/export")
    assert r.status_code == 200
    assert r.headers["content-disposition"].startswith(
        f'attachment; filename="tabel_{cohort["code"]}_'
    )
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    rows = [tuple(c.value for c in row) for row in ws.iter_rows()]
    assert rows[0][:6] == (
        "Студент",
        "Статус",
        "Проведено уроков",
        "Присутствовал",
        "Не отмечено",
        "Посещаемость",
    )
    assert len(rows[0]) == 6 + len(grid["lessons"])
    assert "(отм.)" in rows[0][6 + 2]  # the cancelled lesson is named in its header
    by_name = {row[0]: row for row in rows[1:]}
    cell = ws["F2"]  # a numeric percentage cell, not text
    assert isinstance(cell.value, float) and cell.number_format == "0.0%"
    assert abs(by_name["Алиев"][5] - 1 / 8) < 1e-9
    for s in grid["students"]:
        row = by_name[s["student_name"]]
        assert (row[2], row[3], row[4]) == (s["held"], s["present"], s["unmarked"])
    assert by_name["Алиев"][6:9] == ("П", "Н (болел)", None)  # unmarked cells stay empty
    assert by_name["Бобоев"][6] == "П"
    # an export is audited like every other one
    assert (await sql("SELECT count(*) FROM audit_log WHERE action = 'export'"))[0][0] == 1


async def test_tabel_export_follows_scope(admin: httpx.AsyncClient, login_as: SignedIn) -> None:
    cohort, _ = await past_cohort(admin)
    ins = await login_as("ins", "instructor", branch="Душанбе")  # not their cohort
    r = await ins.get(f"{C}/{cohort['id']}/attendance/export")
    assert (r.status_code, r.json()["code"]) == (404, "not_found")
