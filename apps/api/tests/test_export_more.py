"""Excel of the last gaps: grades tab, student card (two sheets), «Сегодня»/«Не заполнено»,
password reset requests. Same rows as the screen, scope, rights, neutralising, audit entry."""

from io import BytesIO
from typing import Any

import httpx
from openpyxl import load_workbook

from app.export import XLSX
from tests.conftest import SignedIn, make_role, make_user, sql
from tests.edu import C, make_student
from tests.money import enroll as enroll_paid
from tests.money import pay, running_cohort
from tests.test_attendance import enroll, mark, past_cohort
from tests.test_reset_request import _ask


def sheets(r: httpx.Response) -> dict[str, list[tuple[Any, ...]]]:
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    wb = load_workbook(BytesIO(r.content))
    return {ws.title: [tuple(c.value for c in row) for row in ws.iter_rows()] for ws in wb}


async def _exports(entity: str) -> list[Any]:
    return await sql(
        "SELECT user_id, diff FROM audit_log WHERE action = 'export' AND entity = :e ORDER BY id",
        e=entity,
    )


# ---- grades --------------------------------------------------------------------------------


async def test_grades_export_equals_list(admin: httpx.AsyncClient) -> None:
    cohort, lessons = await past_cohort(admin)
    zed, amy = await enroll(admin, cohort, "Zed"), await enroll(admin, cohort, "Amy")
    await mark(admin, lessons[0], amy)
    await admin.patch(f"/api/v1/enrollments/{amy}/grade", json={"homework": 30, "final": 25})
    listed = (await admin.get(f"{C}/{cohort['id']}/grades")).json()
    book = sheets(await admin.get(f"{C}/{cohort['id']}/grades/export"))
    assert list(book) == ["Оценки"]
    rows = book["Оценки"]
    assert [r[0] for r in rows[1:]] == [g["student_name"] for g in listed] == ["Amy", "Zed"]
    assert rows[0][2:6] == ("Д/З", "Финал", "Баллы за посещаемость", "Итого")
    assert [r[5] for r in rows[1:]] == [g["total"] for g in listed]
    assert rows[1][2:4] == (30, 25)
    assert zed
    entry = (await _exports("grade"))[0]
    assert entry[1] == {"rows": 2, "filters": {"cohort": cohort["code"]}}


async def test_grades_export_scope_and_rights(admin: httpx.AsyncClient, login_as: SignedIn) -> None:
    cohort, _ = await past_cohort(admin)
    await enroll(admin, cohort, "Amy")
    istar = await login_as("istar", "manager", branch="Истаравшан")
    assert (await istar.get(f"{C}/{cohort['id']}/grades/export")).status_code == 404
    mine = await login_as("teach", "instructor", branch="Душанбе")  # not their cohort
    assert (await mine.get(f"{C}/{cohort['id']}/grades/export")).status_code == 404
    nobody = await login_as("nobody", await make_role("Пусто", {"cohorts": "view"}))
    assert (await nobody.get(f"{C}/{cohort['id']}/grades/export")).status_code == 403


# ---- student card --------------------------------------------------------------------------


async def test_student_card_has_both_sheets_equal_to_screen(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    enrollment = await enroll_paid(admin, cohort, "=Evil Name")
    student = (await admin.get("/api/v1/students")).json()["items"][0]
    await pay(admin, enrollment, month=1, amount=300)
    await pay(admin, enrollment, month=2, amount=200, confirm=True)
    detail = (await admin.get(f"/api/v1/students/{student['id']}")).json()
    paid = (
        await admin.get("/api/v1/payments", params={"enrollment_id": enrollment, "page_size": 50})
    ).json()["items"]
    book = sheets(await admin.get(f"/api/v1/students/{student['id']}/export"))
    assert list(book) == ["Зачисления", "Оплаты"]
    assert book["Зачисления"][0][:2] == ("Студент", "Филиал")
    assert book["Зачисления"][1][:2] == ("'=Evil Name", "Душанбе")
    assert [r[2] for r in book["Зачисления"][1:]] == [
        e["cohort_code"] for e in detail["enrollments"]
    ]
    assert [r[7] for r in book["Оплаты"][1:]] == [float(p["amount"]) for p in paid]
    assert [r[1] for r in book["Оплаты"][1:]] == ["'=Evil Name"] * 2  # formula neutralised
    entry = (await _exports("student_card"))[0]
    assert entry[1]["rows"] == len(detail["enrollments"]) + len(paid)


async def test_student_card_payments_only_with_payments_right(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    cohort = await running_cohort(admin)
    enrollment = await enroll_paid(admin, cohort, "Amy")
    await pay(admin, enrollment)
    sid = (await admin.get("/api/v1/students")).json()["items"][0]["id"]
    role = await make_role("Только студенты", {"students": "view", "cohorts": "view"})
    viewer = await login_as("viewer", role, branch="Душанбе")
    assert list(sheets(await viewer.get(f"/api/v1/students/{sid}/export"))) == ["Зачисления"]
    nobody = await login_as("nobody", await make_role("Пусто", {"cohorts": "view"}))
    assert (await nobody.get(f"/api/v1/students/{sid}/export")).status_code == 403
    istar = await login_as("istar", "manager", branch="Истаравшан")
    assert (await istar.get(f"/api/v1/students/{sid}/export")).status_code == 404
    other = await make_student("Nobody Else")
    assert (await istar.get(f"/api/v1/students/{other}/export")).status_code == 404


# ---- «Сегодня» / «Не заполнено» ------------------------------------------------------------


async def test_today_export_equals_screen_for_manager(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    cohort, lessons = await past_cohort(admin)
    await enroll(admin, cohort, "Amy")
    await mark(
        admin,
        lessons[-1],
        (await admin.get(f"{C}/{cohort['id']}/grades")).json()[0]["enrollment_id"],
    )
    mgr = await login_as("mgr", "manager", branch="Душанбе")
    screen = (await mgr.get("/api/v1/me/today")).json()
    book = sheets(await mgr.get("/api/v1/me/today/export"))
    assert list(book) == ["Сегодня", "Не заполнено"]
    assert len(book["Сегодня"]) - 1 == len(screen["today"])
    assert len(book["Не заполнено"]) - 1 == len(screen["unfilled"]) > 0
    assert {r[3] for r in book["Не заполнено"][1:]} == {cohort["code"]}
    assert (await _exports("today"))[0][1]["rows"] == len(screen["today"]) + len(screen["unfilled"])
    # another branch sees none of it
    istar = await login_as("istar", "manager", branch="Истаравшан")
    assert len(sheets(await istar.get("/api/v1/me/today/export"))["Не заполнено"]) == 1


async def test_today_export_needs_attendance_right(login_as: SignedIn) -> None:
    nobody = await login_as("nobody", await make_role("Пусто", {"cohorts": "view"}))
    assert (await nobody.get("/api/v1/me/today/export")).status_code == 403


# ---- password reset requests ---------------------------------------------------------------


async def test_reset_requests_export_equals_list_and_scope(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    small = await make_role("Малая", {"users": "view"})
    await make_user("ali", small, branch="Душанбе")
    await make_user("bob", small, branch="Истаравшан")
    await _ask("ali", "198.51.100.1")
    await _ask("bob", "198.51.100.2")
    listed = (await admin.get("/api/v1/users/reset-requests")).json()["items"]
    rows = sheets(await admin.get("/api/v1/users/reset-requests/export"))["Запросы пароля"]
    assert [r[0] for r in rows[1:]] == [i["login"] for i in listed]
    assert sorted(r[0] for r in rows[1:]) == ["ali", "bob"]
    hr = await login_as("hr", await make_role("Кадры", {"users": "change"}), branch="Душанбе")
    scoped = sheets(await hr.get("/api/v1/users/reset-requests/export"))["Запросы пароля"]
    assert [r[0] for r in scoped[1:]] == ["ali"]
    viewer = await login_as("viewer", await make_role("Смотрящий", {"users": "view"}))
    assert (await viewer.get("/api/v1/users/reset-requests/export")).status_code == 403
    assert len(await _exports("reset_request")) == 2  # the 403 wrote nothing


async def test_new_exports_read_as_sentences(admin: httpx.AsyncClient) -> None:
    cohort, _ = await past_cohort(admin)
    await admin.get(f"{C}/{cohort['id']}/grades/export")
    await admin.get("/api/v1/me/today/export")
    items = (await admin.get("/api/v1/audit", params={"action": "export"})).json()["items"]
    texts = [i["text"] for i in items]
    assert any("оценки потока" in t for t in texts)
    assert any("сегодня и не заполнено" in t for t in texts)
