"""Stage 2: payments (#5), idempotent create, confirm flow, debt math, scope, versions, export."""

import asyncio
import uuid
from datetime import date
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError

from app.audit.sentences import Entry, sentence
from app.cohorts.status import today
from tests.conftest import SignedIn, make_role, sql
from tests.money import (
    P,
    add_months,
    cohort_money,
    enroll,
    manager_role,
    money_of,
    pay,
    running_cohort,
    this_month,
)


def key() -> str:
    return str(uuid.uuid4())


async def _key_of(payment_id: str) -> str:
    rows = await sql("SELECT idempotency_key FROM payment WHERE id = :i", i=payment_id)
    return str(rows[0][0])


async def test_create_is_pending_then_confirm_and_unconfirm(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort)
    payment = await pay(admin, e, amount="450.50", note="за октябрь")
    assert (payment["status"], payment["confirmed_by"], payment["amount"]) == (
        "ожидает",
        None,
        "450.50",
    )
    summary = (await admin.get(P)).json()["summary"]
    assert (summary["confirmed"], summary["pending"], summary["count"]) == ("0.00", "450.50", 1)

    r = await admin.post(f"{P}/{payment['id']}/confirm")
    assert r.status_code == 200
    assert r.json()["status"] == "подтверждено" and r.json()["confirmed_by_name"] == "User root"
    assert r.json()["confirmed_at"] is not None
    summary = (await admin.get(P)).json()["summary"]
    assert (summary["confirmed"], summary["pending"]) == ("450.50", "0.00")
    # confirming twice changes nothing and writes no second audit entry
    await admin.post(f"{P}/{payment['id']}/confirm")
    audited = await sql("SELECT count(*) FROM audit_log WHERE action = 'payment.confirm'")
    assert audited[0][0] == 1

    r = await admin.post(f"{P}/{payment['id']}/unconfirm")
    assert r.json()["status"] == "ожидает" and r.json()["confirmed_at"] is None


async def test_create_is_idempotent_on_the_client_key(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort)
    k = uuid.uuid4()
    body = {
        "idempotency_key": str(k),
        "enrollment_id": e,
        "month_no": 1,
        "amount": "500",
        "method": "alif",
        "paid_at": today().isoformat(),
    }
    first = await admin.post(P, json=body)
    retry = await admin.post(P, json=body)
    assert (first.status_code, retry.status_code) == (201, 200)
    assert retry.json()["id"] == first.json()["id"]
    assert (await sql("SELECT count(*) FROM payment"))[0][0] == 1
    assert (await sql("SELECT count(*) FROM audit_log WHERE action = 'payment.create'"))[0][0] == 1
    # the same key for a different payment is refused, not silently merged
    other = await admin.post(P, json={**body, "month_no": 2})
    assert (other.status_code, other.json()["code"]) == (409, "idempotency_key_reused")


async def test_parallel_retries_make_one_payment(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort)
    body = {
        "idempotency_key": key(),
        "enrollment_id": e,
        "month_no": 1,
        "amount": "500",
        "method": "payvand",
        "paid_at": today().isoformat(),
    }
    results = await asyncio.gather(*[admin.post(P, json=body) for _ in range(5)])
    assert sorted(r.status_code for r in results) == [200, 200, 200, 200, 201]
    assert len({r.json()["id"] for r in results}) == 1
    assert (await sql("SELECT count(*) FROM payment"))[0][0] == 1


async def test_bug_05_payment_always_belongs_to_an_enrollment(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort)
    # the database refuses an orphan payment, whatever the code does
    with pytest.raises(IntegrityError):
        await sql(
            "INSERT INTO payment (enrollment_id, month_no, amount, method, paid_at, "
            "idempotency_key) VALUES (NULL, 1, 100, 'alif', :d, gen_random_uuid())",
            d=today(),
        )
    r = await admin.post(
        P,
        json={
            "idempotency_key": key(),
            "enrollment_id": str(uuid.uuid4()),
            "month_no": 1,
            "amount": "100",
            "method": "alif",
            "paid_at": today().isoformat(),
        },
    )
    assert r.status_code == 404
    r = await admin.post(
        P,
        json={
            "idempotency_key": key(),
            "month_no": 1,
            "amount": "100",
            "method": "alif",
            "paid_at": today().isoformat(),
        },
    )
    assert r.status_code == 422
    # and an enrollment that has money cannot be removed from under it
    await pay(admin, e)
    r = await admin.delete(f"/api/v1/enrollments/{e}")
    assert (r.status_code, r.json()["code"]) == (409, "enrollment_has_payments")


async def test_bug_05_payers_never_exceed_enrolled(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    a = await enroll(admin, cohort, "Алиев")
    b = await enroll(admin, cohort, "Бобоев")
    c = await enroll(admin, cohort, "Саидов")
    await pay(admin, a, month=1, confirm=True)
    await pay(admin, a, month=2, confirm=True)  # two payments of one student: one payer
    await pay(admin, b, month=1)  # pending: not a payer yet
    await pay(admin, c, month=1, confirm=True)
    await admin.patch(f"/api/v1/enrollments/{c}", json={"status": "отчислен"})
    m = await cohort_money(admin, cohort)
    # enrolled = not expelled (2), payers = enrolled with a confirmed payment (1: Алиев)
    assert (m["enrolled"], m["payers"]) == (2, 1)
    assert m["payers"] <= m["enrolled"]
    pending = (await admin.get(P, params={"q": "Бобоев"})).json()["items"][0]
    await admin.post(f"{P}/{pending['id']}/confirm")
    m = await cohort_money(admin, cohort)
    assert (m["enrolled"], m["payers"]) == (2, 2)


async def test_payment_values_are_validated(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort)
    base = {
        "idempotency_key": key(),
        "enrollment_id": e,
        "month_no": 1,
        "amount": "100",
        "method": "alif",
        "paid_at": today().isoformat(),
    }
    for bad in (
        {"amount": "0"},
        {"amount": "-5"},
        {"amount": "10.999"},
        {"month_no": 0},
        {"method": "visa"},
        {"paid_at": "2019-12-31"},
        {"paid_at": "2036-01-01"},
        {"idempotency_key": "nope"},
    ):
        r = await admin.post(P, json={**base, **bad})
        assert r.status_code == 422, bad
    r = await admin.post(P, json={**base, "month_no": cohort["months"] + 1})
    assert (r.status_code, r.json()["code"]) == (422, "payment_month_invalid")
    assert r.json()["params"]["months"] == cohort["months"]
    ok = await admin.post(P, json={**base, "month_no": cohort["months"]})
    assert ok.status_code == 201
    # the same limits stand in the database
    insert = (
        "INSERT INTO payment (enrollment_id, month_no, amount, method, paid_at, idempotency_key) "
        "VALUES (:e, :n, :a, :m, :d, gen_random_uuid())"
    )
    for n, amount, method in ((0, 100, "alif"), (1, 0, "alif"), (1, 100, "visa")):
        with pytest.raises(IntegrityError):
            await sql(insert, e=e, n=n, a=amount, m=method, d=today())


async def test_debt_math_mid_course_enrollment_and_partial_payments(
    admin: httpx.AsyncClient,
) -> None:
    cohort = await running_cohort(admin, back=3, weeks=30, price=500)  # started 3 months ago
    full = await enroll(admin, cohort, "С начала")  # months started: 4 (price 500) = 2000
    mid = await enroll(admin, cohort, "Пришёл позже", since=add_months(this_month(), -1))
    # mid: this month and last month = 2 months = 1000
    await pay(admin, full, month=1, amount=500, confirm=True)
    await pay(admin, full, month=2, amount=300, confirm=True)  # partial month 2
    await pay(admin, full, month=2, amount=200, confirm=True)  # the rest, a second payment
    await pay(admin, full, month=3, amount=500)  # pending: does not count
    await pay(admin, mid, month=1, amount=700, confirm=True)
    m = await cohort_money(admin, cohort)
    a, b = money_of(m, full), money_of(m, mid)
    assert (a["expected"], a["paid"], a["debt"], a["paid_months"]) == (
        "2000.00",
        "1000.00",
        "1000.00",
        [1, 2],
    )
    assert (b["expected"], b["paid"], b["debt"]) == ("1000.00", "700.00", "300.00")
    # overpaying never makes a negative debt
    await pay(admin, mid, month=2, amount=900, confirm=True)
    assert money_of(await cohort_money(admin, cohort), mid)["debt"] == "0.00"


async def test_expected_is_capped_at_the_cohort_months_and_starts_with_the_course(
    admin: httpx.AsyncClient,
) -> None:
    # long-running cohort: 6 calendar months, 2 of them behind us, 4 ahead
    cohort = await running_cohort(admin, back=1, weeks=22, price=100)
    e = await enroll(admin, cohort)
    # all months of the cohort the student has "started" so far = this month and last = 2
    assert money_of(await cohort_money(admin, cohort), e)["expected"] == "200.00"
    # a cohort that has not started owes nothing yet
    future = await running_cohort(admin, name="F", back=-1, weeks=8, price=100)
    f = await enroll(admin, future, "Ранний")
    assert money_of(await cohort_money(admin, future), f)["expected"] == "0.00"
    # finished long ago: capped at the cohort's months
    await sql(
        "UPDATE cohort SET start_date = :s, end_date = :e WHERE id = :i",
        s=date(2025, 1, 1),
        e=date(2025, 3, 31),
        i=cohort["id"],
    )
    await sql("UPDATE enrollment SET enrolled_at = :d WHERE id = :i", d=date(2025, 1, 5), i=e)
    assert money_of(await cohort_money(admin, cohort), e)["expected"] == "300.00"
    # an expelled student owes nothing more
    await admin.patch(f"/api/v1/enrollments/{e}", json={"status": "отчислен"})
    assert money_of(await cohort_money(admin, cohort), e)["debt"] == "0.00"


async def test_debtors_list_equals_the_debt_math_and_exports(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin, back=1, weeks=20, price=500)
    owes = await enroll(admin, cohort, "Должник")  # 2 months = 1000
    paid = await enroll(admin, cohort, "Платит")
    await pay(admin, paid, month=1, amount=500, confirm=True)
    await pay(admin, paid, month=2, amount=500, confirm=True)
    await pay(admin, owes, month=1, amount=400, confirm=True)
    r = await admin.get(f"{P}/debtors")
    body = r.json()
    assert [i["student_name"] for i in body["items"]] == ["Должник"]
    assert (body["items"][0]["debt"], body["total_debt"], body["total"]) == (
        "600.00",
        "600.00",
        1,
    )
    x = await admin.get(f"{P}/debtors/export")
    ws = load_workbook(BytesIO(x.content)).active
    assert ws is not None
    rows = [tuple(c.value for c in row) for row in ws.iter_rows()]
    assert rows[0][0] == "Студент" and rows[1][0] == "Должник" and rows[1][8] == 600
    audit = await sql("SELECT diff FROM audit_log WHERE action = 'export' AND entity = 'debtor'")
    assert audit[0][0]["rows"] == 1


async def test_payments_list_filters_totals_and_export_match(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    a, b = await enroll(admin, cohort, "Алиев"), await enroll(admin, cohort, "Бобоев")
    await pay(admin, a, amount=500, method="alif", confirm=True)
    await pay(admin, b, amount=300, method="наличные")
    await pay(admin, b, month=2, amount=200, method="payvand", paid_at=add_months(today(), -1))

    all_ = (await admin.get(P)).json()
    assert all_["total"] == 3 and all_["summary"]["count"] == 3
    only = (await admin.get(P, params={"status": "ожидает"})).json()
    assert (only["total"], only["summary"]["pending"], only["summary"]["confirmed"]) == (
        2,
        "500.00",
        "0.00",
    )
    assert (await admin.get(P, params={"method": "наличные"})).json()["total"] == 1
    month = today().strftime("%Y-%m")
    assert (await admin.get(P, params={"month": month})).json()["total"] == 2
    assert (await admin.get(P, params={"q": "алиев"})).json()["total"] == 1
    assert (await admin.get(P, params={"status": "bad"})).status_code == 422

    params = {"status": "ожидает"}
    listed = (await admin.get(P, params={**params, "page_size": 100})).json()["items"]
    x = await admin.get(f"{P}/export", params=params)
    ws = load_workbook(BytesIO(x.content)).active
    assert ws is not None
    assert [row[1].value for row in ws.iter_rows(min_row=2)] == [i["student_name"] for i in listed]
    # page 2 totals keep describing the whole filter, not the page
    page = (await admin.get(P, params={"page_size": 1, "page": 2})).json()
    assert len(page["items"]) == 1 and page["summary"]["count"] == 3


async def test_options_for_the_payment_form(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin, price=450)
    e = await enroll(admin, cohort, "Алиев")
    await pay(admin, e, month=1, confirm=True)
    r = await admin.get(f"{P}/options", params={"q": "Алиев"})
    (opt,) = r.json()
    assert (opt["enrollment_id"], opt["price_month"], opt["next_month"], opt["months"]) == (
        e,
        "450.00",
        2,
        cohort["months"],
    )
    await admin.patch(f"/api/v1/enrollments/{e}", json={"status": "отчислен"})
    assert (await admin.get(f"{P}/options", params={"q": "Алиев"})).json() == []


async def test_confirmed_payment_is_frozen_and_pending_can_change(
    admin: httpx.AsyncClient,
) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort)
    p = await pay(admin, e, amount=500)
    r = await admin.patch(f"{P}/{p['id']}", json={"amount": "550", "note": "поправка"})
    assert (
        r.status_code == 200 and r.json()["amount"] == "550.00" and r.json()["note"] == "поправка"
    )
    r = await admin.patch(f"{P}/{p['id']}", json={"month_no": cohort["months"] + 1})
    assert r.json()["code"] == "payment_month_invalid"
    await admin.post(f"{P}/{p['id']}/confirm")
    for r in (
        await admin.patch(f"{P}/{p['id']}", json={"amount": "1"}),
        await admin.delete(f"{P}/{p['id']}"),
    ):
        assert (r.status_code, r.json()["code"]) == (409, "payment_confirmed")
    await admin.post(f"{P}/{p['id']}/unconfirm")
    assert (await admin.delete(f"{P}/{p['id']}")).status_code == 204
    assert (await sql("SELECT count(*) FROM payment"))[0][0] == 0


async def test_version_conflict_on_payment_edit(
    login_as: SignedIn, admin: httpx.AsyncClient
) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort)
    p = await pay(admin, e)
    other = await login_as("anna", await manager_role())
    assert (await other.patch(f"{P}/{p['id']}", json={"amount": "510"})).status_code == 200
    r = await admin.patch(f"{P}/{p['id']}", json={"amount": "520"}, headers={"If-Match": "1"})
    assert (r.status_code, r.json()["code"]) == (409, "stale_version")
    assert r.json()["params"]["version"] == 2
    assert float(r.json()["params"]["current"]["amount"]) == 510
    r = await admin.patch(f"{P}/{p['id']}", json={"amount": "520"}, no_version=True)
    assert r.status_code == 428


async def test_branch_scope(login_as: SignedIn, admin: httpx.AsyncClient) -> None:
    dushanbe = await running_cohort(admin, name="D")
    istar = await running_cohort(admin, name="I", branch="Истаравшан")
    ed, ei = (
        await enroll(admin, dushanbe, "Душанбинец"),
        await enroll(admin, istar, "Истаравшанец", branch="Истаравшан"),
    )
    pd, pi = await pay(admin, ed), await pay(admin, ei)
    anna = await login_as("anna", await manager_role(), branch="Душанбе")
    ids = [i["id"] for i in (await anna.get(P)).json()["items"]]
    assert ids == [pd["id"]]
    assert (await anna.get(f"{P}/{pi['id']}")).status_code == 404
    assert (await anna.post(f"{P}/{pi['id']}/confirm")).status_code == 404
    assert (await anna.patch(f"{P}/{pi['id']}", json={"amount": "1"})).status_code == 404
    assert (await anna.delete(f"{P}/{pi['id']}")).status_code == 404
    assert (await anna.get(f"{P}/by-cohort/{istar['id']}")).status_code == 404
    r = await anna.post(
        P,
        json={
            "idempotency_key": key(),
            "enrollment_id": ei,
            "month_no": 1,
            "amount": "1",
            "method": "alif",
            "paid_at": today().isoformat(),
        },
    )
    assert r.status_code == 404
    assert (await anna.get(f"{P}/options", params={"q": "Истаравшанец"})).json() == []
    assert (await anna.get(f"{P}/debtors")).json()["total"] == 1  # only her own branch's debtor
    # a retry of somebody else's key does not leak the payment either
    r = await anna.post(
        P,
        json={
            "idempotency_key": await _key_of(pi["id"]),
            "enrollment_id": ei,
            "month_no": 1,
            "amount": "500",
            "method": "alif",
            "paid_at": today().isoformat(),
        },
    )
    assert r.status_code == 404


async def test_rights_instructor_and_read_only_roles(
    login_as: SignedIn, admin: httpx.AsyncClient
) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort)
    p = await pay(admin, e)
    ins = await login_as("ins", "instructor")
    for r in (
        await ins.get(P),
        await ins.get(f"{P}/{p['id']}"),
        await ins.get(f"{P}/debtors"),
        await ins.get(f"{P}/export"),
        await ins.post(f"{P}/{p['id']}/confirm"),
    ):
        assert r.status_code == 403
    viewer = await login_as("view", await make_role("Бухгалтер", {"payments": "view"}))
    assert (await viewer.get(P)).status_code == 200
    assert (await viewer.post(f"{P}/{p['id']}/confirm")).status_code == 403
    assert (await viewer.delete(f"{P}/{p['id']}")).status_code == 403
    nobody = await login_as("none", await make_role("Никто", {"students": "view"}))
    assert (await nobody.get(P)).status_code == 403
    assert (await nobody.get(f"{P}/export")).status_code == 403


async def test_audit_reads_as_russian_sentences(admin: httpx.AsyncClient) -> None:
    cohort = await running_cohort(admin)
    e = await enroll(admin, cohort, "Алиев Алишер")
    p = await pay(admin, e, amount=500, confirm=True)
    await admin.post(f"{P}/{p['id']}/unconfirm")
    rows = await sql(
        "SELECT action, entity, label, diff FROM audit_log WHERE entity = 'payment' ORDER BY id"
    )
    texts = [sentence(Entry("Анна", a, en, lab, d)) for a, en, lab, d in rows]
    code = cohort["code"]
    assert texts[0].startswith(f"Анна: создана оплата Алиев Алишер ({code}), мес. 1 — ")
    assert "сумма — → 500 TJS" in texts[0] or "сумма" in texts[0]
    assert texts[1] == f"Анна: подтверждена оплата Алиев Алишер ({code}), мес. 1"
    assert texts[2] == f"Анна: снято подтверждение оплаты Алиев Алишер ({code}), мес. 1"
    r = await admin.get("/api/v1/audit", params={"section": "payments"})
    assert r.status_code == 200 and r.json()["total"] == 3
