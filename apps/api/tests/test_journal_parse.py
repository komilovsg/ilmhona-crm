"""The journal reader: header row, name column, year inference, name matching (no database use)."""

import csv
import struct
import time
import zipfile
from datetime import date
from io import BytesIO, StringIO
from typing import Any

import pytest
from openpyxl import Workbook

from app.attendance import journal as jr
from app.core.errors import DomainError
from tests.journal_fixture import journal_xlsx

D = [f"{d:02d}.09" for d in (1, 3, 5)]


def hdr(cell: Any) -> Any:
    out = jr.header_date(cell)
    assert out is not None
    return out


def test_finds_the_date_row_and_the_name_column_under_labels_and_extra_columns() -> None:
    data = journal_xlsx(D, [("Каримов Али", ["б", "н", "у"], False)])
    j = jr.read_journal(data)
    # title, «сс / У1..» row, then the dates row; «№» column before the names
    assert (j.header_row, j.name_col, [c for c, _ in j.dates]) == (2, 1, [2, 3, 4])
    rows = jr.journal_rows(j)
    assert [(r.name, r.marks) for r in rows] == [("Каримов Али", {2: "б", 3: "н", 4: "у"})]


def test_works_without_title_numbers_row_or_number_column_and_with_typed_dates() -> None:
    data = journal_xlsx(
        [date(2025, 9, 1), date(2025, 9, 3)],
        [("Рахимова Мадина", ["б", None], False)],
        title=False,
        numbers=False,
        number_col=False,
        extras=False,
    )
    j = jr.read_journal(data)
    assert (j.header_row, j.name_col) == (0, 0)
    assert [h[:3] for _, h in j.dates] == [(1, 9, 2025), (3, 9, 2025)]


def test_explicit_sheet_row_and_column_override_the_guess() -> None:
    wb = Workbook()
    ws: Any = wb.active
    ws.title = "Пусто"
    ws["A1"] = "ничего"
    data = journal_xlsx(D, [("Каримов Али", ["б", None, None], False)], wb=wb, sheet="Группа")
    j = jr.read_journal(data)  # the right sheet is found on its own
    assert j.table.name == "Группа"
    j = jr.read_journal(data, sheet="2", row=3, col="B")
    assert (j.header_row, j.name_col, j.explicit) == (2, 1, True)
    with pytest.raises(DomainError) as e:
        jr.read_journal(data, sheet="Нет такого")
    assert e.value.error.code == "validation_failed"
    with pytest.raises(DomainError) as e:
        jr.read_journal(data, sheet="Группа", row=1)
    assert e.value.error.code == "journal_no_dates"


def test_csv_with_semicolons_and_dd_mm_text() -> None:
    out = StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["№", "ФИО студента", "01.09", "03.09", "Телефон"])
    w.writerow([1, "Каримов Али", "б", "н", "+992900000001"])
    j = jr.read_journal(out.getvalue().encode("utf-8-sig"))
    assert (j.format, j.header_row, j.name_col) == ("csv", 0, 1)
    assert jr.journal_rows(j)[0].marks == {2: "б", 3: "н"}


def test_not_a_journal_says_what_is_missing() -> None:
    with pytest.raises(DomainError) as e:
        jr.read_journal("ФИО;Телефон\nАли;1".encode())
    assert e.value.error.code == "journal_no_dates"
    with pytest.raises(DomainError) as e:
        jr.read_journal(b"01.09;03.09\n1;2")
    assert e.value.error.code == "journal_no_names"


def test_marks() -> None:
    assert [jr.mark_kind(m) for m in ("б", "Б", "н", "Н.", "у", "x", "=1+1")] == [
        jr.PRESENT,
        jr.PRESENT,
        jr.ABSENT,
        jr.ABSENT,
        jr.EXCUSED,
        None,
        None,
    ]


def test_year_comes_from_the_cohort_dates() -> None:
    start, end = date(2025, 11, 10), date(2026, 2, 20)
    assert jr.resolve_date(hdr("15.12"), start, end) == date(2025, 12, 15)
    assert jr.resolve_date(hdr("12.01"), start, end) == date(2026, 1, 12)
    # a typed date whose default year is wrong (Sheets stamps the year it was typed in)
    assert jr.resolve_date(hdr(date(2024, 12, 15)), start, end) == date(2025, 12, 15)
    # an explicit year in the text wins, even outside the cohort
    assert jr.resolve_date(hdr("15.12.2024"), start, end) == date(2024, 12, 15)
    # outside the cohort without a year: the nearest one
    assert jr.resolve_date(hdr("05.11"), start, end) == date(2025, 11, 5)
    assert jr.resolve_date(hdr("31.02.2025"), start, end) is None
    assert jr.header_date("У1") is None and jr.header_date("1.5 б") is None


@pytest.mark.parametrize(
    ("journal", "crm", "tier"),
    [
        ("каримов ёлкин", "Каримов Елкин", "exact"),
        ("Али Каримов", "Каримов Али", "translit"),
        ("Madina Rahimova", "Рахимова Мадина", "translit"),
        ("Каримов Али Рустамович", "Каримов Али", "tokens"),
        ("Ali Karimof", "Каримов Али", "similar"),
    ],
)
def test_names_match_across_case_yo_script_and_order(journal: str, crm: str, tier: str) -> None:
    other = jr.Person("x", "Каримова Азиза")
    m = jr.match_names([journal], [jr.Person("a", crm), other])[0]
    assert m.person is not None and m.person.key == "a" and m.tier == tier


def test_siblings_and_ambiguity_are_never_guessed() -> None:
    people = [jr.Person("a", "Каримов Али"), jr.Person("b", "Каримова Азиза")]
    assert jr.match_names(["Каримов Азиз"], people)[0].person is None  # other first name
    twins = [jr.Person("a", "Каримов Али Рустамович"), jr.Person("b", "Каримов Али Умаров")]
    m = jr.match_names(["Каримов Али"], twins)[0]
    assert m.person is None and len(m.candidates) == 2
    # two journal rows for one student: the better spelling takes it
    ms = jr.match_names(["Каримов Али", "Али Каримов"], [people[0]])
    assert [x.person is not None for x in ms] == [True, False]


def test_unreadable_and_oversized() -> None:
    with pytest.raises(DomainError) as e:
        jr.read_journal(b"PK\x03\x04 not really a zip")
    assert e.value.error.code == "import_unreadable"
    with pytest.raises(DomainError) as e:
        jr.read_journal(b"\xd0\xcf\x11\xe0 old xls")
    assert e.value.error.code == "import_unreadable"
    with pytest.raises(DomainError) as e:
        jr.read_journal(b"  \n")
    assert e.value.error.code == "import_no_rows"
    rows = "\n".join(f"Имя Фамилия{i};б" for i in range(jr.MAX_ROWS + 5))
    with pytest.raises(DomainError) as e:
        jr.read_journal(f"ФИО;01.09\n{rows}".encode())
    assert e.value.error.code == "import_too_many_rows"
    wb = Workbook()
    ws: Any = wb.active
    for i in range(1, jr.MAX_ROWS + 10):
        ws.cell(i, 1, f"Имя Фамилия{i}")
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(DomainError) as e:
        jr.read_journal(buf.getvalue())
    assert e.value.error.code == "import_too_many_rows"


# ---- hostile files ---------------------------------------------------------------------------


def _small_journal(merge: str | None = None) -> bytes:
    wb = Workbook()
    ws: Any = wb.active
    ws.append(["ФИО", *D])
    ws.append(["Каримов Али", "б", "н", "у"])
    if merge:
        ws.merge_cells(merge)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rezip(data: bytes, patch: dict[str, bytes], *, lie: bool = False) -> bytes:
    """Rewrite some parts of an xlsx; `lie` claims 1000 bytes for every patched part in the zip
    headers (a zip bomb that hides its real size)."""
    src = zipfile.ZipFile(BytesIO(data))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for info in src.infolist():
            z.writestr(info.filename, patch.get(info.filename, src.read(info.filename)))
    raw = bytearray(out.getvalue())
    if lie:
        for name in patch:
            at = 0
            while (at := raw.find(name.encode(), at + 1)) > 0:
                if raw[at - 30 : at - 26] == b"PK\x03\x04":
                    struct.pack_into("<I", raw, at - 30 + 22, 1000)
                if raw[at - 46 : at - 42] == b"PK\x01\x02":
                    struct.pack_into("<I", raw, at - 46 + 24, 1000)
    return bytes(raw)


def test_zip_bomb_with_lying_sizes_is_refused_before_the_parser() -> None:
    data = _small_journal()
    styles = zipfile.ZipFile(BytesIO(data)).read("xl/styles.xml")
    k = styles.index(b">") + 1
    padded = styles[:k] + b" " * (jr.MAX_UNZIPPED + 1) + styles[k:]
    bomb = _rezip(data, {"xl/styles.xml": padded}, lie=True)
    assert len(bomb) < jr.MAX_BYTES
    with pytest.raises(DomainError) as e:
        jr._check_zip(bomb)  # the real size is counted, not the declared one
    assert e.value.error.code == "import_unreadable"


def test_merged_cell_bomb_is_ignored_and_cheap() -> None:
    data = _small_journal("E1:ZZ5000")  # ~3.5 million cells in one merge, a 6 KB file
    started = time.monotonic()
    j = jr.read_journal(data)
    assert time.monotonic() - started < 3
    assert [r.name for r in jr.journal_rows(j)] == ["Каримов Али"]


_LAUGHS = "".join(
    f'<!ENTITY {b} "{("&" + a + ";") * 10}">' for a, b in zip("abcdefg", "bcdefgh", strict=True)
)


@pytest.mark.parametrize(
    "doctype",
    [
        f'<!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">{_LAUGHS}]>',  # billion laughs
        '<!DOCTYPE x [<!ENTITY h SYSTEM "file:///etc/passwd">]>',  # external entity
    ],
)
def test_xml_entities_are_never_expanded(doctype: str) -> None:
    data = _small_journal()
    sheet = zipfile.ZipFile(BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    assert sheet.startswith("<worksheet")  # no xml declaration: the doctype goes first
    cell = '<row r="100"><c r="A100" t="inlineStr"><is><t>&h;</t></is></c></row>'
    sheet = sheet.replace("</sheetData>", f"{cell}</sheetData>", 1)
    evil = _rezip(data, {"xl/worksheets/sheet1.xml": f"{doctype}{sheet}".encode()})
    try:
        j = jr.read_journal(evil)
    except DomainError as e:
        assert e.error.code == "import_unreadable"
        return
    text = repr(j.table.rows)
    assert "root:" not in text and "aaaaaaaaaa" not in text
