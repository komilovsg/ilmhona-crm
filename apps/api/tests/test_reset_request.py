"""«Забыли пароль?»: the same answer for every login, limits per address and login, scope."""

import asyncio
import time
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook

from app.auth import reset
from tests.conftest import SignedIn, make_client, make_role, make_user, sql

A = "/api/v1/auth"
FORGOT = f"{A}/forgot-password"


async def _ask(login: str, ip: str = "203.0.113.7") -> httpx.Response:
    async with make_client(ip) as c:
        return await c.post(FORGOT, json={"login": login})


async def test_same_answer_for_known_unknown_and_inactive() -> None:
    await make_user("anna", "manager")
    await make_user("gone", "manager", active=False)
    answers = [await _ask(x, f"198.51.100.{n}") for n, x in enumerate(("anna", "nobody", "gone"))]
    assert {r.status_code for r in answers} == {202}
    assert len({r.text for r in answers}) == 1
    assert "администратор получит запрос" in answers[0].text
    assert len({tuple(sorted(r.headers)) for r in answers}) == 1


async def test_known_login_is_linked_unknown_is_not() -> None:
    uid = await make_user("anna", "manager")
    await _ask("anna")
    await _ask("nobody", "198.51.100.1")
    rows = await sql("SELECT login, user_id, status, ip FROM password_reset_request ORDER BY login")
    assert [(r[0], r[1], r[2]) for r in rows] == [("anna", uid, "новый"), ("nobody", None, "новый")]
    assert rows[0][3] == "203.0.113.7"


async def test_work_done_is_the_same_for_unknown_login() -> None:
    """A real and an invented login take the same path: one row, one audit entry."""
    await make_user("anna", "manager")
    await _ask("anna", "198.51.100.1")
    await _ask("ghost", "198.51.100.2")
    rows = await sql("SELECT entity_id FROM audit_log WHERE action = 'auth.reset_requested'")
    assert sorted(r[0] for r in rows) == ["anna", "ghost"]


async def test_timing_does_not_separate_known_from_unknown() -> None:
    await make_user("anna", "manager")

    async def mean(login: str, base: int) -> float:
        total = 0.0
        for n in range(6):
            t = time.perf_counter()
            await _ask(login, f"192.0.2.{base + n}")
            total += time.perf_counter() - t
        return total / 6

    known, unknown = await mean("anna", 10), await mean("zzz-unknown", 100)
    assert abs(known - unknown) < 0.15  # no hashing or lookup on only one branch


async def test_limit_per_login_and_still_same_answer() -> None:
    await make_user("anna", "manager")
    answers = [await _ask("anna", f"198.51.100.{n}") for n in range(5)]
    assert {r.status_code for r in answers} == {202}
    assert len({r.text for r in answers}) == 1
    assert (await sql("SELECT count(*) FROM password_reset_request WHERE login = 'anna'"))[0][
        0
    ] == 3
    assert (await _ask("boris", "198.51.100.200")).status_code == 202  # not anna's limit
    assert (await sql("SELECT count(*) FROM password_reset_request WHERE login = 'boris'"))[0][
        0
    ] == 1


async def test_limit_per_address() -> None:
    for n in range(5):
        await _ask(f"user{n}", "203.0.113.50")
    assert (await sql("SELECT count(*) FROM password_reset_request"))[0][0] == 3
    assert (await _ask("other", "203.0.113.51")).status_code == 202
    assert (await sql("SELECT count(*) FROM password_reset_request"))[0][0] == 4


async def test_limit_window_is_one_hour() -> None:
    await make_user("anna", "manager")
    for n in range(3):
        await _ask("anna", f"198.51.100.{n}")
    await sql("UPDATE password_reset_request SET created_at = now() - interval '61 minutes'")
    await _ask("anna", "198.51.100.77")
    assert (await sql("SELECT count(*) FROM password_reset_request"))[0][0] == 4


async def test_list_issue_and_reject(admin: httpx.AsyncClient) -> None:
    anna = await make_user("anna", "manager")
    await make_user("boris", "manager")
    await _ask("anna", "198.51.100.1")
    await _ask("boris", "198.51.100.2")
    await _ask("nobody", "198.51.100.3")
    listing = (await admin.get("/api/v1/users/reset-requests")).json()
    assert listing["total"] == 2  # unknown logins never show
    by_login = {i["login"]: i for i in listing["items"]}
    assert set(by_login) == {"anna", "boris"}

    r = await admin.post(f"/api/v1/users/{anna}/reset-password")  # «Выдать временный пароль»
    assert r.status_code == 200 and r.json()["temp_password"]
    r = await admin.post(f"/api/v1/users/reset-requests/{by_login['boris']['id']}/reject")
    assert r.status_code == 204

    assert (await admin.get("/api/v1/users/reset-requests")).json()["total"] == 0
    rows = await sql(
        "SELECT login, status, handled_by IS NOT NULL FROM password_reset_request "
        "WHERE user_id IS NOT NULL ORDER BY login"
    )
    assert [tuple(x) for x in rows] == [("anna", "выполнен", True), ("boris", "отклонён", True)]
    again = await admin.post(f"/api/v1/users/reset-requests/{by_login['boris']['id']}/reject")
    assert again.status_code == 404
    actions = {x[0] for x in await sql("SELECT action FROM audit_log")}
    assert {"auth.reset_requested", "user.reset_password", "user.reset_rejected"} <= actions


async def test_list_needs_users_change_and_respects_scope(login_as: SignedIn) -> None:
    small = await make_role("Малая", {"users": "view"})  # nobody above `hr` in rights
    await make_user("ali", small, branch="Душанбе")
    await make_user("bob", small, branch="Истаравшан")
    await make_user("rootless", "superadmin")
    await _ask("ali", "198.51.100.1")
    await _ask("bob", "198.51.100.2")
    await _ask("rootless", "198.51.100.3")

    view = await make_role("Смотрящий", {"users": "view"})
    viewer = await login_as("viewer", view)
    assert (await viewer.get("/api/v1/users/reset-requests")).status_code == 403

    change = await make_role("Кадры", {"users": "change"})
    hr = await login_as("hr", change, branch="Душанбе")
    items = (await hr.get("/api/v1/users/reset-requests")).json()["items"]
    assert [i["login"] for i in items] == [
        "ali"
    ]  # own branch; a superadmin's request is not theirs

    bob_req = (await sql("SELECT id FROM password_reset_request WHERE login = 'bob'"))[0][0]
    r = await hr.post(f"/api/v1/users/reset-requests/{bob_req}/reject")
    assert r.status_code == 404  # another branch looks like it does not exist
    root_req = (await sql("SELECT id FROM password_reset_request WHERE login = 'rootless'"))[0][0]
    assert (await hr.post(f"/api/v1/users/reset-requests/{root_req}/reject")).status_code == 404

    boss = await login_as("boss", "superadmin")
    assert (await boss.get("/api/v1/users/reset-requests")).json()["total"] == 3


async def test_anonymous_cannot_read_requests(client: httpx.AsyncClient) -> None:
    assert (await client.get("/api/v1/users/reset-requests")).status_code == 401


async def test_issued_password_works(admin: httpx.AsyncClient) -> None:
    uid = await make_user("anna", "manager", password="Forgotten1!x")
    await _ask("anna")
    pw = (await admin.post(f"/api/v1/users/{uid}/reset-password")).json()["temp_password"]
    async with make_client() as c:
        r = await c.post(f"{A}/login", json={"login": "anna", "password": pw})
        assert r.status_code == 200 and r.json()["must_change_password"] is True


async def test_global_cap_stops_a_spread_out_flood(monkeypatch: pytest.MonkeyPatch) -> None:
    """Many addresses, many invented logins: the table still stops growing (and the answer
    stays the same)."""
    monkeypatch.setattr(reset, "GLOBAL_LIMIT", 5)
    answers = {(await _ask(f"bot{n}", f"192.0.2.{n + 1}")).text for n in range(8)}
    assert len(answers) == 1
    assert (await sql("SELECT count(*) FROM password_reset_request"))[0][0] == 5


async def test_concurrent_rejects_handle_a_request_once(admin: httpx.AsyncClient) -> None:
    await make_user("anna", "manager")
    await _ask("anna")
    rid = (await sql("SELECT id FROM password_reset_request"))[0][0]
    path = f"/api/v1/users/reset-requests/{rid}/reject"
    answers = await asyncio.gather(*(admin.post(path) for _ in range(4)))
    assert sorted(r.status_code for r in answers) == [204, 404, 404, 404]
    rows = await sql("SELECT count(*) FROM audit_log WHERE action = 'user.reset_rejected'")
    assert rows[0][0] == 1


async def test_attacker_login_text_is_neutralised_in_audit_export(
    admin: httpx.AsyncClient,
) -> None:
    await _ask('=HYPERLINK("http://evil.example","x")')
    r = await admin.get("/api/v1/audit/export?action=auth.reset_requested")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    texts = [str(c.value) for row in ws.iter_rows() for c in row if "HYPERLINK" in str(c.value)]
    assert texts and all(not t.startswith("=") for t in texts)
