"""Stage 3: leads. IDs (#14), name hygiene (#15), duplicates, convert, CSV import, scope, rights."""

import uuid
from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook

from app.audit.sentences import Entry, sentence
from app.leads.hygiene import clean_name, normalise_phone, phone_key
from tests.conftest import SignedIn, branch_id, make_role, sql
from tests.edu import instructor_rights, make_cohort, make_course, make_room, make_student
from tests.sales import L, csv_bytes, make_campaign, make_lead, manager_role

# ---- #14: ids are generated -------------------------------------------------------------


async def test_bug_14_lead_ids_are_generated_by_the_database(admin: httpx.AsyncClient) -> None:
    a, b = await make_lead(admin, "Анна"), await make_lead(admin, "Борис")
    assert uuid.UUID(a["id"]) != uuid.UUID(b["id"])  # no «L-001» counter, no repeats
    # the client cannot pick an id
    r = await admin.post(
        L, json={"id": "L-001", "full_name": "Х", "branch_id": str(await branch_id("Душанбе"))}
    )
    assert r.status_code == 422
    # and a direct insert gets one from the column default
    rows = await sql(
        "INSERT INTO lead (full_name, branch_id) SELECT 'Z', id FROM branch LIMIT 1 RETURNING id"
    )
    assert isinstance(rows[0][0], uuid.UUID)


# ---- #15: name hygiene ------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "name", "review"),
    [
        ("🔥 Алишер 🔥", "Алишер", True),
        ("Нигина ❤️", "Нигина", True),
        ("𝗟𝗶𝘃𝗲 𝗦𝗰𝗵𝗼𝗼𝗹", "Live School", True),
        ("𝗟𝗶𝘃𝗲𝗿𝗽𝗼𝗼𝗹 𝗙𝗮𝗻 𝗘𝗱𝗶𝘁𝘀 🔴", "Liverpool Fan Edits", True),
        ("Ａｌｉ Ｋａｒｉｍｏｖ", "Ali Karimov", True),
        ("Али​бер‍ди", "Алиберди", True),
        ("  Нигина    Рахимова  ", "Нигина Рахимова", False),
        ("✨ Мадина ✨", "Мадина", True),
        ("Ӣсо Нурализода", "Ӣсо Нурализода", False),
        ("Dilnoza O'Neil", "Dilnoza O'Neil", False),
        ("Курбон | Школа Эльбрус", "Курбон Школа Эльбрус", True),
        ("Best Shop > Дешёвые телефоны", "Best Shop Дешёвые телефоны", True),
        ("Магазин www.shop.tj", "Магазин www.shop.tj", True),
        ("https://t.me/ali", "https://t.me/ali", True),
        ("latipov_008_", "latipov_008", True),
        ("@dilya.tj", "dilya.tj", True),
        ("Сухроб 2024", "Сухроб 2024", True),
        ("Анна Мария Фон Дер Ляйен Смит", "Анна Мария Фон Дер Ляйен Смит", True),
        ("✨✨", "Без имени", True),
        ("123", "123", True),
        ("", "Без имени", True),
        ('=HYPERLINK("http://x.y")', 'HYPERLINK("http://x.y', True),
    ],
)
def test_bug_15_name_is_cleaned_and_flagged(raw: str, name: str, review: bool) -> None:
    got = clean_name(raw)
    assert (got.name, got.needs_review) == (name, review)


async def test_bug_15_api_keeps_the_original_and_flags(admin: httpx.AsyncClient) -> None:
    ok = await make_lead(admin, "🔥 Алишер 🔥")
    assert (ok["full_name"], ok["raw_name"], ok["needs_review"]) == (
        "Алишер",
        "🔥 Алишер 🔥",
        True,
    )
    bad = await make_lead(admin, "latipov_008_")
    assert (bad["full_name"], bad["raw_name"], bad["needs_review"]) == (
        "latipov_008",
        "latipov_008_",
        True,
    )
    only = (await admin.get(L, params={"needs_review": "true"})).json()
    assert {i["id"] for i in only["items"]} == {ok["id"], bad["id"]}
    # a person fixes the name: it is re-checked and the flag goes
    r = await admin.patch(f"{L}/{bad['id']}", json={"full_name": "Латипов Фарход"})
    assert (r.json()["full_name"], r.json()["needs_review"]) == ("Латипов Фарход", False)
    # or just confirms it as it is
    again = await make_lead(admin, "Best Shop | Phones")
    r = await admin.patch(f"{L}/{again['id']}", json={"needs_review": False})
    assert r.json()["needs_review"] is False


def test_bug_15_phones_are_normalised() -> None:
    for raw in ("901234567", "+992 90 123 45 67", "992901234567", "00992901234567", "0901234567"):
        assert normalise_phone(raw) == "+992901234567", raw
    assert normalise_phone("12345") is None
    assert normalise_phone("+7 495 123 45 67") is None
    assert phone_key("+992 (90) 123-45-67") == phone_key("901234567") == "901234567"
    assert phone_key("12345") is None


async def test_phone_is_normalised_and_the_raw_one_kept(admin: httpx.AsyncClient) -> None:
    a = await make_lead(admin, "Анна", phone="90 123 45 67")
    assert (a["phone"], a["phone_raw"]) == ("+992901234567", "90 123 45 67")
    b = await make_lead(admin, "Борис", phone="+7 495 123 45 67")
    assert (b["phone"], b["phone_raw"]) == ("+7 495 123 45 67", "+7 495 123 45 67")
    c = await make_lead(admin, "Вера")
    assert (c["phone"], c["phone_raw"]) == (None, None)


# ---- duplicates -------------------------------------------------------------------------


async def test_duplicate_phone_warns_with_links_and_never_blocks(admin: httpx.AsyncClient) -> None:
    first = await make_lead(admin, "Анна", phone="+992 90 000 00 01")
    assert first["duplicates"] is None
    sid = await make_student("Старый студент", phone="900000001")
    second = await make_lead(admin, "Анна вернулась", phone="900 000 001")  # 201: not blocked
    dup = second["duplicates"]
    assert [x["id"] for x in dup["leads"]] == [first["id"]]
    assert [x["id"] for x in dup["students"]] == [str(sid)]
    assert (await admin.get(L)).json()["total"] == 2
    # editing the phone checks again; other fields do not
    third = await make_lead(admin, "Вера", phone="911111111")
    r = await admin.patch(f"{L}/{third['id']}", json={"phone": "900000001"})
    assert len(r.json()["duplicates"]["leads"]) == 2
    assert (
        "duplicates" not in (await admin.patch(f"{L}/{third['id']}", json={"comment": "x"})).json()
    )


async def test_duplicates_stay_inside_the_users_branch(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    await make_lead(admin, "Истаравшанская", phone="900000001", branch="Истаравшан")
    await make_student("Студент ИСТ", branch="Истаравшан", phone="900000001")
    anna = await login_as("anna", await manager_role(), branch="Душанбе")
    r = await make_lead(anna, "Анна", phone="900000001")
    assert r["duplicates"] is None  # nothing about another branch's people leaks


# ---- convert ----------------------------------------------------------------------------


async def _cohort(admin: httpx.AsyncClient, branch: str = "Душанбе") -> dict[str, Any]:
    tag = uuid.uuid4().hex[:6]
    return await make_cohort(
        admin, await make_course(f"Python {tag}"), await make_room(branch, tag), branch=branch
    )


async def test_convert_creates_a_student_and_is_idempotent(admin: httpx.AsyncClient) -> None:
    lead = await make_lead(admin, "Анна Каримова", phone="901112233", status="нулевой урок")
    r = await admin.post(f"{L}/{lead['id']}/convert", json={})
    assert r.status_code == 200, r.text
    body = r.json()
    sid = body["student"]["id"]
    assert (body["student"]["created"], body["lead"]["status"]) == (True, "записан")
    assert body["lead"]["converted_student_id"] == sid
    student = (await admin.get(f"/api/v1/students/{sid}")).json()
    assert (student["full_name"], student["phone"]) == ("Анна Каримова", "+992901112233")
    assert (await sql("SELECT source_lead_id::text FROM student"))[0][0] == lead["id"]
    # again: the same student, nothing new, no second audit entry
    audit = await sql("SELECT count(*) FROM audit_log WHERE action LIKE 'lead.convert'")
    again = (await admin.post(f"{L}/{lead['id']}/convert", json={})).json()
    assert (again["student"]["id"], again["student"]["created"], again["already"]) == (
        sid,
        False,
        True,
    )
    assert (await sql("SELECT count(*) FROM student"))[0][0] == 1
    assert await sql("SELECT count(*) FROM audit_log WHERE action LIKE 'lead.convert'") == audit


async def test_convert_links_the_existing_student_with_the_same_phone(
    admin: httpx.AsyncClient,
) -> None:
    old = await make_student("Анна Старая", phone="+992 90 111 22 33")
    other_branch = await make_student("Из Истаравшана", branch="Истаравшан", phone="901112233")
    lead = await make_lead(admin, "Анна Новая", phone="901112233")
    body = (await admin.post(f"{L}/{lead['id']}/convert", json={})).json()
    assert (body["student"]["id"], body["student"]["created"]) == (str(old), False)
    assert body["student"]["id"] != str(other_branch)  # same branch only
    assert (await sql("SELECT count(*) FROM student"))[0][0] == 2
    # a lead from the other branch with the same phone links to its own branch's student
    lead2 = await make_lead(admin, "Другая", phone="901112233", branch="Истаравшан")
    body2 = (await admin.post(f"{L}/{lead2['id']}/convert", json={})).json()
    assert body2["student"]["id"] == str(other_branch)


async def test_convert_enrolls_into_a_cohort_once(admin: httpx.AsyncClient) -> None:
    cohort = await _cohort(admin)
    lead = await make_lead(admin, "Анна", phone="901112233")
    url = f"{L}/{lead['id']}/convert"
    body = (await admin.post(url, json={"cohort_id": cohort["id"]})).json()
    assert body["enrollment"]["cohort_code"] == cohort["code"]
    assert body["enrollment"]["new"] is True
    again = (await admin.post(url, json={"cohort_id": cohort["id"]})).json()
    assert (again["enrollment"]["id"], again["enrollment"]["new"]) == (
        body["enrollment"]["id"],
        False,
    )
    assert (await sql("SELECT count(*) FROM enrollment"))[0][0] == 1
    # converted first, enrolled later: the second call enrolls
    lead2 = await make_lead(admin, "Борис", phone="902223344")
    await admin.post(f"{L}/{lead2['id']}/convert", json={})
    later = (
        await admin.post(f"{L}/{lead2['id']}/convert", json={"cohort_id": cohort["id"]})
    ).json()
    assert later["enrollment"]["new"] is True
    assert (await sql("SELECT count(*) FROM enrollment"))[0][0] == 2


async def test_convert_cohort_rules(admin: httpx.AsyncClient, login_as: SignedIn) -> None:
    lead = await make_lead(admin, "Анна", phone="901112233")
    url = f"{L}/{lead['id']}/convert"
    foreign = await _cohort(admin, "Истаравшан")
    r = await admin.post(url, json={"cohort_id": foreign["id"]})
    assert (r.status_code, r.json()["code"]) == (422, "student_branch_mismatch")
    assert (await sql("SELECT count(*) FROM student"))[0][0] == 0  # nothing half-done
    r = await admin.post(url, json={"cohort_id": str(uuid.uuid4())})
    assert r.status_code == 404
    # a role that edits leads but not cohorts cannot enroll
    role = await make_role("Лиды", {"leads": "change", "students": "change"})
    pat = await login_as("pat", role)
    cohort = await _cohort(admin)
    r = await pat.post(url, json={"cohort_id": cohort["id"]})
    assert r.status_code == 403
    assert (await pat.post(url, json={})).status_code == 200


async def test_status_rules_around_conversion(admin: httpx.AsyncClient) -> None:
    lead = await make_lead(admin, "Анна", phone="901112233")
    r = await admin.patch(f"{L}/{lead['id']}", json={"status": "записан"})
    assert (r.status_code, r.json()["code"]) == (422, "lead_use_convert")
    await admin.post(f"{L}/{lead['id']}/convert", json={})
    r = await admin.patch(f"{L}/{lead['id']}", json={"status": "отказ"})
    assert (r.status_code, r.json()["code"]) == (409, "lead_converted")
    # other fields still change
    assert (await admin.patch(f"{L}/{lead['id']}", json={"comment": "ок"})).status_code == 200


# ---- call status, callbacks, funnel -----------------------------------------------------


async def test_callback_today_list_and_default_time(admin: httpx.AsyncClient) -> None:
    now = datetime.now(UTC)
    due = await make_lead(
        admin, "Сегодня", call_status="перезвонить", next_contact_at=now.isoformat()
    )
    await make_lead(
        admin, "Просрочен", call_status="перезвонить",
        next_contact_at=(now - timedelta(days=2)).isoformat(),
    )  # fmt: skip
    await make_lead(
        admin, "Завтра", call_status="перезвонить",
        next_contact_at=(now + timedelta(days=2)).isoformat(),
    )  # fmt: skip
    await make_lead(admin, "Позвонили", call_status="позвонили", next_contact_at=now.isoformat())
    names = [
        i["full_name"] for i in (await admin.get(L, params={"callback": "true"})).json()["items"]
    ]
    assert names == ["Просрочен", "Сегодня"]  # oldest first
    summary = (await admin.get("/api/v1/metrics/leads")).json()
    assert summary["callbacks_today"] == 2
    # picking «перезвонить» with no time means now: it lands on today's list
    quick = await make_lead(admin, "Быстрый")
    r = await admin.patch(f"{L}/{quick['id']}", json={"call_status": "перезвонить"})
    assert r.json()["next_contact_at"] is not None
    assert (await admin.get(L, params={"callback": "true"})).json()["total"] == 3
    assert due["next_contact_at"] is not None


async def test_funnel_counts_and_conversion(admin: httpx.AsyncClient) -> None:
    for i, status in enumerate(("новый", "новый", "заинтересован", "нулевой урок", "отказ")):
        await make_lead(admin, f"Лид {i}", status=status, phone=f"90000000{i}")
    converted = await make_lead(admin, "Купил", phone="911111111")
    await admin.post(f"{L}/{converted['id']}/convert", json={})
    f = (await admin.get("/api/v1/metrics/leads")).json()
    assert f["counts"] == {
        "новый": 2,
        "заинтересован": 1,
        "нулевой урок": 1,
        "записан": 1,
        "отказ": 1,
        "всего": 6,
    }
    assert f["conversion"] == 16.7
    empty = (
        await admin.get(
            "/api/v1/metrics/leads", params={"date_from": "2020-01-01", "date_to": "2020-01-02"}
        )
    ).json()
    assert (empty["counts"]["всего"], empty["conversion"]) == (0, 0.0)
    # the period is the day the lead was created
    await sql("UPDATE lead SET created_at = '2026-01-15 10:00+05' WHERE full_name = 'Лид 0'")
    jan = (
        await admin.get(
            "/api/v1/metrics/leads", params={"date_from": "2026-01-15", "date_to": "2026-01-15"}
        )
    ).json()
    assert jan["counts"]["всего"] == 1


async def test_list_filters_and_search(admin: httpx.AsyncClient) -> None:
    course = await make_course("Python")
    camp = await make_campaign(admin)
    await make_lead(admin, "Алишер", phone="901000001", source="Instagram", course_id=str(course))
    await make_lead(admin, "Борис", phone="902000002", campaign_id=camp["id"])
    await make_lead(admin, "Вера", branch="Истаравшан", call_status="не ответил")

    async def names(**params: str) -> set[str]:
        r = await admin.get(L, params=params)
        return {i["full_name"] for i in r.json()["items"]}

    assert await names(source="Instagram") == {"Алишер"}
    assert await names(course_id=str(course)) == {"Алишер"}
    assert await names(campaign_id=camp["id"]) == {"Борис"}
    assert await names(call_status="не ответил") == {"Вера"}
    assert await names(branch_id=str(await branch_id("Истаравшан"))) == {"Вера"}
    assert await names(q="902 000") == {"Борис"}
    assert await names(q="алиш") == {"Алишер"}
    assert await names(status="отказ") == set()
    assert (await admin.get(L, params={"status": "bogus"})).status_code == 422


async def test_version_conflict_on_a_lead(admin: httpx.AsyncClient) -> None:
    lead = await make_lead(admin, "Анна")
    ok = await admin.patch(f"{L}/{lead['id']}", json={"comment": "a"}, headers={"If-Match": "1"})
    assert ok.status_code == 200
    stale = await admin.patch(f"{L}/{lead['id']}", json={"comment": "b"}, headers={"If-Match": "1"})
    assert (stale.status_code, stale.json()["code"]) == (409, "stale_version")
    assert (
        await admin.patch(f"{L}/{lead['id']}", json={"comment": "c"}, no_version=True)
    ).status_code == 428


async def test_delete_and_refusal_reason(admin: httpx.AsyncClient) -> None:
    lead = await make_lead(admin, "Анна", status="отказ", refusal_reason="дорого")
    assert lead["refusal_reason"] == "дорого"
    r = await admin.patch(f"{L}/{lead['id']}", json={"status": "заинтересован"})
    assert r.json()["refusal_reason"] is None  # a reason belongs to a refusal only
    assert (await admin.delete(f"{L}/{lead['id']}")).status_code == 204
    assert (await admin.get(f"{L}/{lead['id']}")).status_code == 404


# ---- CSV import -------------------------------------------------------------------------

FB = [
    ["id", "created_time", "campaign_name", "platform", "full_name", "email", "phone_number"],
    ["l:1", "2026-10-01", "Осень", "fb", "🔥 Алишер Каримов", "a@x.tj", "+992901112233"],
    ["l:2", "2026-10-01", "Осень", "ig", "latipov_008_", "", "902223344"],
    ["l:3", "2026-10-02", "Осень", "fb", "Нигина | Школа", "", "903334455"],
    ["l:4", "2026-10-02", "Осень", "fb", "", "", ""],
    ["l:5", "2026-10-03", "Осень", "fb", "Дубль в файле", "", "+992 90 111 22 33"],
]


async def _post(c: httpx.AsyncClient, path: str, data: bytes, **params: str) -> httpx.Response:
    return await c.post(
        f"{L}{path}", content=data, params=params, headers={"Content-Type": "text/csv"}
    )


async def _counts() -> tuple[int, int]:
    return (
        (await sql("SELECT count(*) FROM lead"))[0][0],
        (await sql("SELECT count(*) FROM audit_log WHERE entity = 'lead'"))[0][0],
    )


async def test_import_preview_writes_nothing_and_flags(admin: httpx.AsyncClient) -> None:
    await make_student("Старый", phone="902223344")
    before = await _counts()
    r = await _post(
        admin, "/import/preview", csv_bytes(FB), branch_id=str(await branch_id("Душанбе"))
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert await _counts() == before  # a preview is a read
    assert body["mapping"]["full_name"] == "full_name"
    assert body["mapping"]["phone"] == "phone_number"
    assert body["mapping"]["source"] == "platform"
    rows = {x["line"]: x for x in body["rows"]}
    assert (rows[2]["full_name"], rows[2]["needs_review"], rows[2]["source"]) == (
        "Алишер Каримов",
        True,
        "Facebook",
    )
    assert (rows[2]["phone"], rows[2]["raw_name"]) == ("+992901112233", "🔥 Алишер Каримов")
    assert (rows[3]["needs_review"], rows[3]["source"], rows[3]["phone"]) == (
        True,
        "Instagram",
        "+992902223344",
    )
    assert [d["kind"] for d in rows[3]["duplicates"]] == ["student"]  # known student
    assert rows[4]["needs_review"] is True  # «Нигина | Школа»
    assert rows[5]["skip"] == "empty"
    assert [d["kind"] for d in rows[6]["duplicates"]] == ["file"]  # same phone as line 2
    assert body["totals"] == {
        "rows": 5,
        "importable": 4,
        "skipped": 1,
        "needs_review": 3,
        "duplicates": 2,
    }


async def test_import_confirm_creates_leads(admin: httpx.AsyncClient) -> None:
    camp = await make_campaign(admin)
    course = await make_course("Python")
    r = await _post(
        admin,
        "/import",
        csv_bytes(FB),
        branch_id=str(await branch_id("Душанбе")),
        campaign_id=camp["id"],
        course_id=str(course),
    )
    assert r.status_code == 201, r.text
    assert (r.json()["created"], r.json()["skipped"], r.json()["needs_review"]) == (4, 1, 3)
    rows = (await admin.get(L)).json()["items"]
    assert len(rows) == 4
    assert {i["campaign_id"] for i in rows} == {camp["id"]}
    assert {i["course_name"] for i in rows} == {"Python"}
    assert {i["raw_name"] for i in rows} >= {"🔥 Алишер Каримов", "latipov_008_"}
    log = await sql("SELECT action FROM audit_log WHERE entity = 'lead'")
    assert [x[0] for x in log] == ["lead.import"]
    audit = (await admin.get("/api/v1/audit", params={"entity": "lead"})).json()["items"]
    assert "импортировано лидов из CSV: 4, из них на проверку: 3" in audit[0]["text"]


async def test_import_mapping_semicolons_utf16_and_cp1251(admin: httpx.AsyncClient) -> None:
    branch = str(await branch_id("Душанбе"))
    custom = [["Клиент", "Мобильный", "Куда"], ["Анна", "901112233", "TikTok"]]
    r = await _post(admin, "/import/preview", csv_bytes(custom, ";"), branch_id=branch)
    assert r.json()["mapping"]["full_name"] is None  # nothing to guess from
    import json

    mapping = json.dumps({"full_name": "Клиент", "phone": "Мобильный", "source": "Куда"})
    r = await _post(
        admin, "/import/preview", csv_bytes(custom, ";"), branch_id=branch, mapping=mapping
    )
    row = r.json()["rows"][0]
    assert (row["full_name"], row["phone"], row["source"]) == ("Анна", "+992901112233", "TikTok")
    # a Facebook export is UTF-16 with a tab delimiter
    fb16 = csv_bytes(FB[:2], "\t").decode().encode("utf-16")
    r = await _post(admin, "/import/preview", fb16, branch_id=branch)
    assert r.json()["rows"][0]["full_name"] == "Алишер Каримов"
    ansi = csv_bytes(FB[:2]).decode().encode("cp1251", errors="replace")
    assert (await _post(admin, "/import/preview", ansi, branch_id=branch)).status_code == 200
    bad = json.dumps({"full_name": "нет такой колонки"})
    r = await _post(admin, "/import/preview", csv_bytes(FB), branch_id=branch, mapping=bad)
    assert r.json()["code"] == "import_mapping_invalid"


async def test_import_limits_and_garbage(admin: httpx.AsyncClient) -> None:
    from app.leads import importer

    branch = str(await branch_id("Душанбе"))
    many = [FB[0], *([["x", "", "", "fb", "Анна", "", "901112233"]] * (importer.MAX_ROWS + 1))]
    r = await _post(admin, "/import/preview", csv_bytes(many), branch_id=branch)
    assert (r.status_code, r.json()["code"]) == (422, "import_too_many_rows")
    big = b"full_name,phone\n" + b"a" * (importer.MAX_BYTES + 1)
    r = await _post(admin, "/import/preview", big, branch_id=branch)
    assert (r.status_code, r.json()["code"]) == (413, "import_too_large")
    r = await _post(admin, "/import/preview", b"full_name,phone\n", branch_id=branch)
    assert r.json()["code"] == "import_no_rows"
    r = await _post(admin, "/import/preview", b"\x00\x01\x02\x03", branch_id=branch)
    assert r.status_code == 422
    assert (await _counts())[0] == 0


async def test_import_is_formula_safe(admin: httpx.AsyncClient) -> None:
    table = [
        ["full_name", "phone_number", "comment"],
        ['=HYPERLINK("http://evil.example","click")', "901112233", "=1+1"],
        ["@SUM(A1)", "902223344", "+cmd|' /C calc'!A0"],
    ]
    r = await _post(admin, "/import", csv_bytes(table), branch_id=str(await branch_id("Душанбе")))
    assert r.status_code == 201
    for lead in (await admin.get(L)).json()["items"]:
        assert not (lead["full_name"] or "").startswith(("=", "+", "-", "@"))
        assert not (lead["comment"] or "").startswith(("=", "+", "-", "@"))
    exported = await admin.get(f"{L}/export")
    ws = load_workbook(BytesIO(exported.content)).active
    assert ws is not None
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert not any(isinstance(v, str) and v.startswith(("=", "@")) for v in row)


async def test_import_needs_change_rights(login_as: SignedIn) -> None:
    viewer = await login_as("view", await make_role("Только чтение", {"leads": "view"}))
    r = await _post(viewer, "/import/preview", csv_bytes(FB))
    assert r.status_code == 403


# ---- scope, rights, export --------------------------------------------------------------


async def test_branch_scope(admin: httpx.AsyncClient, login_as: SignedIn) -> None:
    mine = await make_lead(admin, "Душанбинский", branch="Душанбе")
    other = await make_lead(admin, "Истаравшанский", branch="Истаравшан")
    anna = await login_as("anna", await manager_role(), branch="Душанбе")
    ids = {i["id"] for i in (await anna.get(L)).json()["items"]}
    assert ids == {mine["id"]}
    assert (await anna.get(f"{L}/{other['id']}")).status_code == 404
    assert (await anna.patch(f"{L}/{other['id']}", json={"comment": "x"})).status_code == 404
    assert (await anna.delete(f"{L}/{other['id']}")).status_code == 404
    assert (await anna.post(f"{L}/{other['id']}/convert", json={})).status_code == 404
    # her new leads go to her branch; another branch is refused
    ok = await make_lead(anna, "Новый", branch="Душанбе")
    assert ok["branch_name"] == "Душанбе"
    r = await anna.post(
        L, json={"full_name": "Чужой", "branch_id": str(await branch_id("Истаравшан"))}
    )
    assert r.status_code == 403
    assert (await anna.get("/api/v1/metrics/leads")).json()["counts"]["всего"] == 2
    assert (await admin.get("/api/v1/metrics/leads")).json()["counts"]["всего"] == 3
    exported = await anna.get(f"{L}/export")
    ws = load_workbook(BytesIO(exported.content)).active
    assert ws is not None
    assert {r[0] for r in ws.iter_rows(min_row=2, values_only=True)} == {"Душанбинский", "Новый"}


async def test_instructor_has_no_access_even_with_the_right(
    admin: httpx.AsyncClient, login_as: SignedIn
) -> None:
    lead = await make_lead(admin, "Анна")
    async with instructor_rights(leads="change", marketing="change"):
        ins = await login_as("ins", "instructor", branch="Душанбе")
        for method, url in (
            ("GET", L),
            ("GET", f"{L}/{lead['id']}"),
            ("GET", f"{L}/export"),
            ("GET", "/api/v1/metrics/leads"),
            ("POST", f"{L}/{lead['id']}/convert"),
            ("DELETE", f"{L}/{lead['id']}"),
        ):
            r = await ins.request(method, url, json={} if method == "POST" else None)
            assert r.status_code == 403, (method, url)
        assert (await ins.get("/api/v1/campaigns")).status_code == 403


async def test_view_only_cannot_write(admin: httpx.AsyncClient, login_as: SignedIn) -> None:
    lead = await make_lead(admin, "Анна")
    viewer = await login_as("view", await make_role("Только чтение", {"leads": "view"}))
    assert (await viewer.get(L)).status_code == 200
    assert (await viewer.get(f"{L}/export")).status_code == 200
    assert (await viewer.post(L, json={"full_name": "x"})).status_code == 403
    assert (await viewer.patch(f"{L}/{lead['id']}", json={"comment": "x"})).status_code == 403
    assert (await viewer.post(f"{L}/{lead['id']}/convert", json={})).status_code == 403
    nobody = await login_as("nobody", await make_role("Никто", {"students": "view"}))
    assert (await nobody.get(L)).status_code == 403


async def test_export_equals_filtered_list_and_is_audited(admin: httpx.AsyncClient) -> None:
    await make_lead(admin, "Алишер", phone="901000001", source="Instagram")
    await make_lead(admin, "Борис", phone="902000002", comment="=SUM(A1)")
    r = await admin.get(f"{L}/export", params={"q": "о"})
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("ФИО", "Телефон", "Курс")
    listed = (await admin.get(L, params={"q": "о"})).json()["items"]
    assert [r[0] for r in rows[1:]] == [i["full_name"] for i in listed]
    comments = [r[11] for r in rows[1:]]
    assert "'=SUM(A1)" in comments  # neutralised
    audit = await sql("SELECT count(*) FROM audit_log WHERE action = 'export' AND entity = 'lead'")
    assert audit[0][0] == 1


async def test_audit_sentences(admin: httpx.AsyncClient) -> None:
    lead = await make_lead(admin, "Анна Каримова", phone="901112233")
    await admin.patch(f"{L}/{lead['id']}", json={"status": "заинтересован"})
    await admin.post(f"{L}/{lead['id']}/convert", json={})
    texts = [
        i["text"]
        for i in (await admin.get("/api/v1/audit", params={"section": "leads"})).json()["items"]
    ]
    assert any("создан лид Анна Каримова" in t for t in texts)
    assert any("статус новый → заинтересован" in t for t in texts)
    assert any("лид Анна Каримова записан — создан новый студент Анна Каримова" in t for t in texts)
    assert sentence(Entry("Курбон", "lead.delete", "lead", "Анна", {})) == "Курбон: удалён лид Анна"


async def test_implausible_phone_goes_to_review(admin: httpx.AsyncClient) -> None:
    for bad in ("abc", "12", "+7"):
        assert (await make_lead(admin, "Анна Каримова", phone=bad))["needs_review"] is True, bad
    assert (await make_lead(admin, "Анна Каримова", phone="+7 495 123 45 67"))[
        "needs_review"
    ] is False
    assert (await make_lead(admin, "Анна Каримова", phone="901234567"))["needs_review"] is False
    good = await make_lead(admin, "Борис Каримов", phone="901234567")
    r = await admin.patch(f"{L}/{good['id']}", json={"phone": "12"})
    assert r.json()["needs_review"] is True
    table = [
        ["full_name", "phone_number"],
        ["Анна Каримова", "abc"],
        ["Борис Каримов", "901112233"],
    ]
    r = await _post(
        admin, "/import/preview", csv_bytes(table), branch_id=str(await branch_id("Душанбе"))
    )
    assert [x["needs_review"] for x in r.json()["rows"]] == [True, False]
