"""One generic Excel export for every table.

A list module declares its columns (`Col`) and exposes `GET <list>/export` next to the list. The
export reuses the list's filter function and select (the module builds both from the same
`where`), so it can never drift from what the screen shows: same filters, search, sort, scope.

Cells are typed (dates are dates, money is a number with a TJS format), the header row is frozen,
column widths fit the content, text that Excel could read as a formula is neutralised, and every
export writes an `export` audit entry (who, which table, filters, row count).
"""

import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Literal

from fastapi import Response
from fastapi.concurrency import run_in_threadpool
from sqlalchemy.ext.asyncio import AsyncSession

from app.audit.log import record
from app.auth.deps import Principal
from app.cohorts.status import TZ
from app.core.errors import DomainError, errors

MAX_ROWS = 50_000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = '#,##0.00 "TJS"'
EXPORT_TOO_LARGE = errors.define("export_too_large", 422, "Too many rows for one Excel file")

Kind = Literal["text", "int", "date", "datetime", "money", "bool", "percent"]
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_PHONE = re.compile(r"^\+\d[\d ()-]*$")


@dataclass(frozen=True)
class Col[T]:
    header: str
    value: Callable[[T], Any]
    kind: Kind = "text"


def neutralise(text: str) -> str:
    """Text that starts with = + - @ tab or CR is prefixed with ' so Excel never runs it. A plain
    phone number (`+992 ...`) cannot be a formula and keeps its form."""
    if text.startswith(_FORMULA_START) and not _PHONE.match(text):
        return "'" + text
    return text


def _cell(value: Any, kind: Kind) -> tuple[Any, str | None]:
    """(value for the sheet, number format)."""
    if value is None:
        return None, None
    if kind == "bool":
        return ("Да" if value else "Нет"), None
    if kind == "money":
        return float(Decimal(value)), MONEY_FORMAT
    if kind == "percent":  # given as 55.6 (percent), stored as a real percentage cell
        return float(Decimal(value) / 100), "0.0%"
    if kind == "int":
        return int(value), "0"
    if kind == "datetime":
        local = value.astimezone(TZ).replace(tzinfo=None) if value.tzinfo else value
        return local, "dd.mm.yyyy hh:mm"
    if kind == "date":
        if isinstance(value, datetime):  # the Dushanbe day, as the UI shows it, not the UTC one
            value = (value.astimezone(TZ) if value.tzinfo else value).date()
        return value, "dd.mm.yyyy"
    return neutralise(str(value)), None


def _shown_len(value: Any) -> int:
    if isinstance(value, datetime):
        return 16
    if isinstance(value, date):
        return 10
    if isinstance(value, float):
        return len(f"{value:,.2f} TJS")
    return len(str(value)) if value is not None else 0


def build_xlsx[T](title: str, columns: Sequence[Col[T]], rows: Sequence[T]) -> bytes:
    # imported here: openpyxl is slow to import and only exports need it (cold start)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title[:31]
    ws.append([c.header for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    widths = [len(c.header) for c in columns]
    for r, row in enumerate(rows, start=2):
        for i, c in enumerate(columns):
            value, fmt = _cell(c.value(row), c.kind)
            if value is None:
                continue
            cell = ws.cell(row=r, column=i + 1, value=value)
            if fmt:
                cell.number_format = fmt
            widths[i] = max(widths[i], _shown_len(value))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 60)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def check_cap(rows: Sequence[Any]) -> None:
    """Callers fetch at most MAX_ROWS + 1 rows; one extra means the table is too big."""
    if len(rows) > MAX_ROWS:
        raise DomainError(EXPORT_TOO_LARGE, limit=MAX_ROWS)


async def xlsx_response[T](
    db: AsyncSession,
    p: Principal,
    *,
    entity: str,
    slug: str,
    title: str,
    columns: Sequence[Col[T]],
    rows: Sequence[T],
    filters: dict[str, Any],
) -> Response:
    """Build the file and write the audit entry (same transaction as the request)."""
    check_cap(rows)
    # CPU-bound: off the event loop, so one big export does not stall every other request
    data = await run_in_threadpool(build_xlsx, title, columns, rows)
    await record(
        db,
        p.id,
        "export",
        entity,
        None,
        {"rows": len(rows), "filters": {k: v for k, v in filters.items() if v not in (None, "")}},
        branch_id=p.branch_id,
    )
    name = f"{slug}_{datetime.now(UTC).astimezone(TZ):%Y-%m-%d}.xlsx"
    return Response(
        data,
        media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
