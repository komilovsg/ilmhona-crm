"""A teacher's «Журнал класса» as a table, and the rules that read it. No database here.

The journal is a Google Sheet: a header row of lesson dates (date cells or «dd.mm» text, maybe a
row above with «У1…У24» lesson numbers and a «сс» / «№» label column), a column of student names
(sometimes with «№» before it), marks «б» (был), «н» (не был), «у» (уважительная причина), an empty
cell = no lesson / no mark. A struck-through name = the student left the group. Other columns
(«Телефон», «Telegram», «✅ Посещено», «📊 % посещ.») are ignored.

Reading is formula-safe: a file is only ever read (`data_only`, cached values), never calculated,
and nothing from it is stored except the marks we match to enrollments.
"""

import csv
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any

from app.core.errors import VALIDATION_FAILED, DomainError, errors
from app.importer.parse import norm_name, same_person, translit_key
from app.leads.errors import IMPORT_NO_ROWS, IMPORT_TOO_MANY_ROWS, IMPORT_UNREADABLE

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLS = 200
MAX_SHEETS = 12
MAX_UNZIPPED = 40 * 1024 * 1024
HEADER_SCAN_ROWS = 40
STRIKE_COLS = 8  # names sit in the first columns: strikes are looked up there only
MAX_TEXT = 200

JOURNAL_NO_DATES = errors.define("journal_no_dates", 422, "No row with lesson dates found")
JOURNAL_NO_NAMES = errors.define("journal_no_names", 422, "No column with student names found")

Cell = str | date | None

PRESENT, ABSENT, EXCUSED = "present", "absent", "excused"
EXCUSED_REASON = "уважительная причина"
_MARKS = {
    "б": PRESENT,
    "был": PRESENT,
    "была": PRESENT,
    "+": PRESENT,
    "н": ABSENT,
    "нб": ABSENT,
    "не был": ABSENT,
    "не была": ABSENT,
    "у": EXCUSED,
    "уп": EXCUSED,
}
_EMPTY_MARKS = {"", "-", "–", "—", ".", "·"}
_DATE = re.compile(r"^(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2}|\d{4}))?\.?$")
_NAME_LABEL = re.compile(r"фио|ф\.\s?и\.\s?о|студент|ученик|слушател|имя|name", re.I)
_SERVICE_ROW = re.compile(r"^(итого|всего|итог|средн|количество|кол-во|процент|%|посещ)", re.I)
_WORD = re.compile(r"[^\W\d_]{2,}(?:[-'’][^\W\d_]+)*")


@dataclass
class Table:
    name: str
    rows: list[list[Cell]]
    struck: set[tuple[int, int]] = field(default_factory=set)  # (row, col), 0-based


@dataclass
class Journal:
    """What was read from the file: its sheets and the sheet that was picked."""

    format: str
    sheets: list[str]
    table: Table
    header_row: int  # 0-based
    name_col: int  # 0-based
    dates: list[tuple[int, tuple[int, int, int | None, bool]]]  # (col, (day, month, year, strong))
    explicit: bool


# ---- reading the file -----------------------------------------------------------------------


def _norm_cell(v: object) -> Cell:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int | float):
        return str(int(v)) if float(v).is_integer() else str(v)
    if not isinstance(v, str):
        return None  # times, durations: not a journal value
    s = re.sub(r"\s+", " ", v.replace("\xa0", " ")).strip()
    if not s or (s.startswith("#") and s.endswith(("!", "?", "A"))):  # #REF! #N/A #NAME?
        return None
    return s[:MAX_TEXT]


def _trim(rows: list[list[Cell]]) -> list[list[Cell]]:
    while rows and all(c is None for c in rows[-1]):
        rows.pop()
    return rows


def _read_csv(data: bytes) -> Table:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DomainError(IMPORT_UNREADABLE)
    first = text.split("\n", 1)[0]
    delimiter = max(",;\t", key=first.count)
    rows: list[list[Cell]] = []
    try:
        for record in csv.reader(StringIO(text), delimiter=delimiter):
            if len(rows) > MAX_ROWS:
                raise DomainError(IMPORT_TOO_MANY_ROWS, limit=MAX_ROWS)
            rows.append([_norm_cell(c) for c in record[:MAX_COLS]])
    except csv.Error:
        raise DomainError(IMPORT_UNREADABLE) from None
    return Table("CSV", _trim(rows))


def _check_zip(data: bytes) -> None:
    """Refuse a zip bomb before the XML parser sees it."""
    try:
        with zipfile.ZipFile(BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > 400 or sum(i.file_size for i in infos) > MAX_UNZIPPED:
                raise DomainError(IMPORT_UNREADABLE)
    except zipfile.BadZipFile:
        raise DomainError(IMPORT_UNREADABLE) from None


def _read_xlsx(data: bytes, only: str | None) -> tuple[list[str], list[Table]]:
    from openpyxl import load_workbook  # slow to import: only an import needs it

    _check_zip(data)
    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=False)
    except Exception:  # noqa: BLE001 - any parser failure is «not a readable file»
        raise DomainError(IMPORT_UNREADABLE) from None
    names = list(wb.sheetnames)[:MAX_SHEETS]
    tables = []
    for name in names:
        if only is not None and name != only:
            continue
        ws = wb[name]
        rows: list[list[Cell]] = []
        struck: set[tuple[int, int]] = set()
        for r, cells in enumerate(
            ws.iter_rows(min_row=1, max_row=MAX_ROWS + 1, max_col=MAX_COLS), start=0
        ):
            row: list[Cell] = []
            for c, cell in enumerate(cells):
                value = _norm_cell(cell.value)
                row.append(value)
                if c < STRIKE_COLS and isinstance(value, str) and cell.font and cell.font.strike:
                    struck.add((r, c))
            rows.append(row)
        rows = _trim(rows)
        if len(rows) > MAX_ROWS:
            raise DomainError(IMPORT_TOO_MANY_ROWS, limit=MAX_ROWS)
        tables.append(Table(name, rows, struck))
    wb.close()
    return names, tables


# ---- header row, name column ----------------------------------------------------------------


def header_date(cell: Cell) -> tuple[int, int, int | None, bool] | None:
    """(day, month, year or None, year is explicit in the text) for a date cell or «dd.mm» text."""
    if isinstance(cell, date):
        return cell.day, cell.month, cell.year, False  # a typed date: its year may be a default
    if not isinstance(cell, str):
        return None
    m = _DATE.match(cell)
    if not m:
        return None
    day, month = int(m.group(1)), int(m.group(2))
    if not (1 <= day <= 31 and 1 <= month <= 12):
        return None
    year = m.group(3)
    if year is None:
        return day, month, None, False
    return day, month, int(year) + 2000 if len(year) == 2 else int(year), True


def resolve_date(hdr: tuple[int, int, int | None, bool], start: date, end: date) -> date | None:
    """The calendar date of a header cell. The year comes from the cohort: the one that puts the
    day inside [start, end], else the nearest. A typed date or a text with a year that already
    lies inside the cohort is kept as written."""
    day, month, year, strong = hdr
    if year is not None:
        try:
            exact = date(year, month, day)
        except ValueError:
            return None
        if strong or start <= exact <= end:
            return exact
    candidates = []
    for y in range(start.year - 1, end.year + 2):
        try:
            candidates.append(date(y, month, day))
        except ValueError:
            continue
    inside = [c for c in candidates if start <= c <= end]
    if inside:
        return inside[0]
    if year is not None:
        try:
            return date(year, month, day)
        except ValueError:
            return None
    if not candidates:
        return None
    return min(candidates, key=lambda c: min(abs((c - start).days), abs((c - end).days)))


def _dates_in(row: list[Cell]) -> list[tuple[int, tuple[int, int, int | None, bool]]]:
    out = []
    for c, cell in enumerate(row):
        hdr = header_date(cell)
        if hdr:
            out.append((c, hdr))
    return out


def find_header_row(rows: list[list[Cell]]) -> int | None:
    """The row (among the first 40) with the most date cells; the first one on a tie."""
    best, best_n = None, 0
    for r, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        n = len(_dates_in(row))
        if n > best_n:
            best, best_n = r, n
    return best


def looks_like_name(cell: Cell) -> bool:
    if not isinstance(cell, str) or cell.startswith(("@", "✅", "📊", "+", "http")):
        return False
    if re.search(r"\d", cell) or cell.casefold().strip(" .") in _MARKS or _SERVICE_ROW.match(cell):
        return False
    return bool(_WORD.search(cell)) and sum(len(w) for w in _WORD.findall(cell)) >= 3


def find_name_col(rows: list[list[Cell]], header: int, skip: set[int]) -> int | None:
    """The column under the header that holds the most names (two words weigh double); a header
    label like «ФИО студента» tips a close call."""
    width = max((len(r) for r in rows), default=0)
    best, best_score = None, 0
    for c in range(width):
        if c in skip:
            continue
        score = 0
        for row in rows[header + 1 :]:
            cell = row[c] if c < len(row) else None
            if looks_like_name(cell):
                score += 2 if len(_WORD.findall(str(cell))) >= 2 else 1
        labels = [rows[r][c] for r in range(max(0, header - 2), header + 1) if c < len(rows[r])]
        if score and any(isinstance(x, str) and _NAME_LABEL.search(x) for x in labels):
            score += 50
        if score > best_score:
            best, best_score = c, score
    return best


def col_index(raw: str) -> int | None:
    """`3` or `C` -> 2 (0-based)."""
    raw = raw.strip()
    if raw.isdigit() and int(raw) >= 1:
        return int(raw) - 1
    if raw.isalpha() and raw.isascii() and len(raw) <= 2:
        n = 0
        for ch in raw.upper():
            n = n * 26 + ord(ch) - 64
        return n - 1
    return None


def _detect(
    table: Table, row: int | None, col: int | None
) -> tuple[int, int, list[tuple[int, tuple[int, int, int | None, bool]]]] | None:
    header = row if row is not None else find_header_row(table.rows)
    if header is None or header >= len(table.rows):
        return None
    dates = _dates_in(table.rows[header])
    if not dates:
        return None
    name_col = col if col is not None else find_name_col(table.rows, header, {c for c, _ in dates})
    if name_col is None:
        return None
    return header, name_col, dates


def read_journal(
    data: bytes, *, sheet: str | None = None, row: int | None = None, col: str | None = None
) -> Journal:
    """Parse the upload (xlsx or csv, told by its first bytes) and find the dates and the names.
    `sheet` (name or 1-based number), `row` (1-based) and `col` (number or letter) override the
    guess. Raises journal_no_dates / journal_no_names when the file does not look like a journal."""
    if not data.strip():
        raise DomainError(IMPORT_NO_ROWS)
    name_col = None
    if col:
        name_col = col_index(col)
        if name_col is None or name_col >= MAX_COLS:
            raise DomainError(VALIDATION_FAILED)
    if row is not None and not 1 <= row <= MAX_ROWS:
        raise DomainError(VALIDATION_FAILED)
    header = row - 1 if row else None

    if data[:2] == b"PK":
        fmt = "xlsx"
        names, tables = _read_xlsx(data, None)
    elif data[:4] == b"\xd0\xcf\x11\xe0":  # the old .xls container
        raise DomainError(IMPORT_UNREADABLE)
    else:
        fmt = "csv"
        table = _read_csv(data)
        names, tables = [table.name], [table]
    if not tables or not any(t.rows for t in tables):
        raise DomainError(IMPORT_NO_ROWS)

    if sheet:
        picked = next((t for t in tables if t.name == sheet), None)
        if picked is None and sheet.isdigit() and 1 <= int(sheet) <= len(tables):
            picked = tables[int(sheet) - 1]
        if picked is None:
            raise DomainError(VALIDATION_FAILED)
        candidates = [picked]
    else:
        candidates = tables
    best: tuple[int, Table, tuple[int, int, list[Any]]] | None = None
    for t in candidates:
        found = _detect(t, header, name_col)
        if found and (best is None or len(found[2]) > best[0]):
            best = (len(found[2]), t, found)
    if best is None:
        # say what is missing: the dates, or the names
        t = candidates[0]
        h = header if header is not None else find_header_row(t.rows)
        raise DomainError(
            JOURNAL_NO_DATES if h is None or not _dates_in(t.rows[h]) else JOURNAL_NO_NAMES
        )
    _, table, (h, nc, dates) = best
    return Journal(fmt, names, table, h, nc, dates, bool(sheet or row or col))


# ---- rows and marks -------------------------------------------------------------------------


@dataclass
class JournalRow:
    line: int  # 1-based sheet row
    name: str
    struck: bool
    marks: dict[int, str]  # col -> the text of a non-empty cell


def journal_rows(j: Journal) -> list[JournalRow]:
    date_cols = [c for c, _ in j.dates]
    out = []
    for r in range(j.header_row + 1, len(j.table.rows)):
        cells = j.table.rows[r]
        name = cells[j.name_col] if j.name_col < len(cells) else None
        if not looks_like_name(name):
            continue
        marks = {}
        for c in date_cols:
            v = cells[c] if c < len(cells) else None
            if isinstance(v, str) and v.casefold().strip(" .") not in _EMPTY_MARKS:
                marks[c] = v
        out.append(JournalRow(r + 1, str(name), (r, j.name_col) in j.table.struck, marks))
    return out


def mark_kind(text: str) -> str | None:
    """present / absent / excused for «б» / «н» / «у» and spelled-out forms; None = unknown."""
    return _MARKS.get(text.casefold().strip(" ."))


# ---- names ----------------------------------------------------------------------------------


@dataclass(frozen=True)
class Person:
    key: Any  # enrollment id
    name: str


@dataclass
class Match:
    person: Person | None
    tier: str | None = None  # exact / translit / tokens / similar
    candidates: tuple[str, ...] = ()  # when ambiguous


def _tokens(name: str) -> list[str]:
    return translit_key(name).split()


def _overlap(a: list[str], b: list[str]) -> int:
    left, n = list(b), 0
    for w in a:
        if w in left:
            left.remove(w)
            n += 1
    return n


def _same_person(a: str, b: str) -> bool:
    return bool(same_person(a, b))


_TIERS: tuple[tuple[str, Callable[[str, str], bool]], ...] = (
    ("exact", lambda a, b: norm_name(a) == norm_name(b)),
    ("translit", lambda a, b: translit_key(a) == translit_key(b)),
    ("tokens", lambda a, b: _overlap(_tokens(a), _tokens(b)) >= 2),
    ("similar", _same_person),
)


def match_names(names: list[str], people: list[Person]) -> list[Match]:
    """One Match per journal name. Tier by tier (exact spelling, same letters in any order or
    script, two equal words, the importer's same-person rule): a name is matched when exactly one
    free person fits and nobody else claims that person at the same tier. Anything left that fit
    several people is ambiguous (never guessed), the rest has no match."""
    result: list[Match | None] = [None] * len(names)
    taken: set[Any] = set()
    ambiguous: dict[int, tuple[str, ...]] = {}
    for tier, fits in _TIERS:
        wants: dict[int, list[Person]] = {}
        for i, n in enumerate(names):
            if result[i] is None:
                wants[i] = [p for p in people if p.key not in taken and fits(n, p.name)]
        claims: dict[Any, int] = {}
        for found in wants.values():
            for p in found:
                claims[p.key] = claims.get(p.key, 0) + 1
        for i, found in wants.items():
            if len(found) == 1 and claims[found[0].key] == 1:
                result[i] = Match(found[0], tier)
                taken.add(found[0].key)
                ambiguous.pop(i, None)
            elif found and i not in ambiguous:
                ambiguous[i] = tuple(p.name for p in found)
    return [r or Match(None, None, ambiguous.get(i, ())) for i, r in enumerate(result)]
