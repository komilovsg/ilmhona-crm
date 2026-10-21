"""The import report. `report-*.md` carries counts and old ids only (it may be shared with the
management); `review-*.xlsx` lists the rows that need a human and DOES hold names and phones, so
both files stay in the gitignored `import-data/`."""

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SHEETS = (
    "COURSES",
    "INSTRUCTORS",
    "COHORTS",
    "STUDENTS",
    "ENROLLMENTS",
    "PAYMENTS",
    "ATTENDANCE",
    "GRADES",
    "MARKETING",
    "LEADS",
    "INSTRUCTOR_PAYMENTS",
)
MAX_IDS = 60  # old ids listed per reason in the md report


@dataclass
class ReviewRow:
    sheet: str
    old_id: str
    reason: str
    details: dict[str, Any]


@dataclass
class Recon:
    label: str
    new: str
    old: str
    comment: str


@dataclass
class Report:
    source: str = ""
    mode: str = "dry-run"
    started: datetime = field(default_factory=datetime.now)
    read: Counter[str] = field(default_factory=Counter)
    imported: Counter[str] = field(default_factory=Counter)  # rows created in the new database
    existing: Counter[str] = field(default_factory=Counter)  # found by legacy_ref: nothing added
    merged: Counter[str] = field(default_factory=Counter)  # old rows folded into another one
    dropped: dict[str, dict[str, list[str]]] = field(
        default_factory=lambda: defaultdict(lambda: defaultdict(list))
    )
    review: list[ReviewRow] = field(default_factory=list)
    notes: Counter[str] = field(default_factory=Counter)
    recon: list[Recon] = field(default_factory=list)
    ledgers: list[tuple[str, list[tuple[str, str, str]]]] = field(default_factory=list)

    def drop(self, sheet: str, reason: str, old_id: str) -> None:
        self.dropped[sheet][reason].append(old_id or "—")

    def flag(self, sheet: str, reason: str, old_id: str, **details: Any) -> None:
        self.review.append(ReviewRow(sheet, old_id or "—", reason, details))

    def note(self, text: str, n: int = 1) -> None:
        self.notes[text] += n

    def dropped_count(self, sheet: str) -> int:
        return sum(len(v) for v in self.dropped.get(sheet, {}).values())

    def review_count(self, sheet: str) -> int:
        return sum(1 for r in self.review if r.sheet == sheet)

    # ---- the md report: counts and old ids only ---------------------------------------------

    def markdown(self) -> str:
        out = [
            "# Перенос из старой системы — отчёт",
            "",
            f"Режим: **{self.mode}** · снимок: `{self.source}` · {self.started:%Y-%m-%d %H:%M}",
            "",
            "## По листам",
            "",
            "| Лист | Прочитано | Перенесено | Уже было | Слито дублей | Отброшено | На проверку |",
            "|---|---:|---:|---:|---:|---:|---:|",
        ]
        for s in SHEETS:
            out.append(
                f"| {s} | {self.read[s]} | {self.imported[s]} | {self.existing[s]} | "
                f"{self.merged[s]} | {self.dropped_count(s)} | {self.review_count(s)} |"
            )
        out += ["", "## Отброшено (причина, старые ID)", ""]
        any_dropped = False
        for s in SHEETS:
            for reason, ids in sorted(self.dropped.get(s, {}).items()):
                any_dropped = True
                shown = ", ".join(ids[:MAX_IDS]) + (" …" if len(ids) > MAX_IDS else "")
                out.append(f"- **{s}** — {reason}: {len(ids)} ({shown})")
        if not any_dropped:
            out.append("Ничего.")
        out += ["", "## На проверку (причина, число, старые ID)", ""]
        by_reason: dict[tuple[str, str], list[str]] = defaultdict(list)
        for r in self.review:
            by_reason[(r.sheet, r.reason)].append(r.old_id)
        if not by_reason:
            out.append("Ничего.")
        for (s, reason), ids in sorted(by_reason.items()):
            shown = ", ".join(ids[:MAX_IDS]) + (" …" if len(ids) > MAX_IDS else "")
            out.append(f"- **{s}** — {reason}: {len(ids)} ({shown})")
        out += ["", "## Что сделано по правилам", ""]
        for text, n in sorted(self.notes.items()):
            out.append(f"- {text}: {n}")
        if not self.notes:
            out.append("—")
        out += [
            "",
            "## Сверка с интерфейсом старой системы",
            "",
            "| Показатель | Новая система (`app/metrics`) | Старая система | Расхождение |",
            "|---|---:|---:|---|",
        ]
        for line in self.recon:
            out.append(f"| {line.label} | {line.new} | {line.old} | {line.comment} |")
        for title, lines in self.ledgers:
            out += ["", f"## {title}", "", "| Статья | Строк | Сумма, TJS |", "|---|---:|---:|"]
            out += [f"| {a} | {b} | {c} |" for a, b, c in lines]
        out.append("")
        return "\n".join(out)

    # ---- the review list: with personal data -------------------------------------------------

    def review_workbook(self) -> Workbook:
        wb = Workbook()
        wb.remove(wb.active)  # type: ignore[arg-type]
        for s in SHEETS:
            rows = [r for r in self.review if r.sheet == s]
            if not rows:
                continue
            ws = wb.create_sheet(s[:31])
            keys: list[str] = []
            for r in rows:
                keys += [k for k in r.details if k not in keys]
            ws.append(["Старый ID", "Причина", *keys])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for r in rows:
                ws.append([r.old_id, r.reason, *[_plain(r.details.get(k)) for k in keys]])
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"  # never a formula, whatever the text starts with
            ws.freeze_panes = "A2"
            for i in range(1, len(keys) + 3):
                ws.column_dimensions[get_column_letter(i)].width = 28
        if not wb.sheetnames:
            wb.create_sheet("Пусто")
        return wb

    def write(self, directory: Path, stamp: str) -> tuple[Path, Path]:
        directory.mkdir(parents=True, exist_ok=True)
        md = directory / f"report-{stamp}.md"
        xlsx = directory / f"review-{stamp}.xlsx"
        md.write_text(self.markdown(), encoding="utf-8")
        md.chmod(0o600)
        self.review_workbook().save(xlsx)
        xlsx.chmod(0o600)
        return md, xlsx


def _plain(value: Any) -> Any:
    if value is None or isinstance(value, int | float | str):
        return value
    return str(value)
